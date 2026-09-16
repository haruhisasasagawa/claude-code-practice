# -*- coding: utf-8 -*-
"""
v2.0(店舗版)のワークブックに「仮数値」を入れた入力済みサンプルを作る。
実データは一切使わず、売上・在庫・原価CSV(期間A/B)とMSO商品CSV(金曜4週分)を
決まった乱数で合成して貼り込み、動員数・保持時間・作るタイミング・ピーク時間・事前準備率も入れる。
  python make_sample_v2.py <ver2.0.xlsx> <出力.xlsx> [--csv-dir DIR] [--sunday 2026-09-13] [--no-label]
    --csv-dir : 合成したCSV(期間A/B・MSO4週)と verify_v2.py 用の引数(CSVは相対パス)を書き出す
                (既定: <出力.xlsx>と同じフォルダの sample_csv/)
    --sunday  : 期間Aの最終日。日曜以外はエラー終了。省略時は今日以前で直近の日曜
    --no-label: 各シートの「【サンプル：数値はすべて仮の値です】」表記を入れない
検証: cd <CSV出力先> && python <repo>/verify_v2.py <再計算済みサンプル.xlsx> $(cat verify_args.txt)
"""
import argparse
import csv
import datetime as dt
import json
import os
import random
import sys
import zlib

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_calib import CALIB_SHEETS, MSO_HEADERS, MSO_MAX, MSO_NCOL, read_mso_rows  # noqa: E402
from build_tool import CSV_HEADERS, CSV_MAX, NCOL, N_SLOTS, read_csv_rows            # noqa: E402
from upgrade_v2 import VIEW_ADDR, VIEW_ALL, check_hidden_cols, lock_workbook, restore_comment_vml               # noqa: E402

SEED = 20260913
ROW_P0 = 14                                   # 期間データ: 商品1行目
ATT_A = [8000, 12000, 11000]                  # 期間A 金・土・日 の動員数(仮)
ATT_B = [7500, 11500, 10500, 5000, 4500, 5500, 5000]   # 期間B 金〜木(仮)
PEAK_ATT = 1200                               # ピーク動員数(仮)
RATE = 80                                     # 事前準備率(%)。100だと販売予測数と作る数が同じになり効果が見えない
PEAK_START, PEAK_END = (18, 30), (19, 30)     # ⑥ピーク時間(仮)
HT_THR = 30                                   # 「直前に作る」の基準(分)
CSV_A_DAYS, CSV_B_DAYS = 3, 7
STALE_DAYS = 9                                # ツール側の「古いデータ」判定: TODAY()−終了日 > 9
TAG = "【サンプル：数値はすべて仮の値です】"
THEATER = ("0000", "サンプル劇場")              # 劇場コード・劇場名(実在コードと取り違えないよう仮の値)

# 店舗商品の仮の設定: (名前に含まれる語, 購買率, 保持時間(分), 1日のMSO明細数, 帯別の売れ方の重み①〜⑤, 単価)
# 名前の部分一致で当てる(店舗版の登録商品名がそのまま使えるように)。該当なしは既定値
PRODUCT_RULES = [
    ("ポップチキン",      0.022, 20,  60, (0.8, 1.5, 1.0, 1.0, 0.6), 500),
    ("トリュフソルトバタ", 0.018, 30,  55, (1.0, 1.1, 1.0, 1.1, 0.7), 450),
    ("シチリアハーブ",    0.012, 30,  30, (1.0, 1.0, 1.0, 1.0, 0.8), 450),
    ("ケチャップ＆マスタード", 0.015, 15, 50, (1.3, 1.2, 0.9, 0.9, 0.5), 550),
    ("４種のチーズ",      0.009, 15,  25, (1.2, 1.1, 0.9, 1.0, 0.5), 600),
    ("ハラペーニョ",      0.006, 15,  18, (0.9, 1.0, 1.0, 1.2, 0.8), 650),
    ("ビーフシチュ",      0.007, 15,  20, (0.7, 1.0, 1.1, 1.3, 0.8), 700),
    ("ジョンソンヴィル",  0.008, 20,  22, (1.0, 1.1, 1.0, 1.0, 0.6), 650),
    ("２００％",          0.005, 20,  15, (0.8, 1.0, 1.1, 1.2, 0.7), 750),
    ("シナモンシュガー",  0.014, 60,  40, (0.5, 0.8, 1.3, 1.5, 0.7), 400),
    ("チョコクリーム",    0.009, 45,  25, (0.5, 0.8, 1.3, 1.4, 0.8), 450),
    ("レインボー",        0.004, 45,  12, (0.4, 0.9, 1.4, 1.3, 0.6), 450),
    ("レモン＆ラムネ",    0.003, 45,   5, (0.4, 0.9, 1.4, 1.3, 0.6), 450),   # 4週で30個未満 → 少データ
]
DEFAULT_RULE = (None, 0.008, 30, 20, (1.0, 1.0, 1.0, 1.0, 1.0), 500)
MANUAL_PRIORITY = {"シナモンシュガー": "直前"}    # 手動の例(保持時間は長いが、売れ筋なので焼きたてを出す)
MEMOS = {"ポップチキン": "仮のメモ：揚げ時間3分",
         "シナモンシュガー": "仮のメモ：売れ筋なので焼きたてを出したい → 手動で「直前」",
         "レモン＆ラムネ": "仮のメモ：MSO4週の合計が30個未満なので全体の時間帯係数で計算（商品別の波シート参照）"}

# 売上CSVの店舗商品以外の行(小分類, 商品名, 購買率, 単価)。除外小分類はプルダウンに出ない
FILLER_SALES = [
    ("ポップコーン", "ポップコーン　塩Ｍ", 0.045, 500), ("ポップコーン", "ポップコーン　キャラメルＭ", 0.055, 500),
    ("ポップコーン", "ハーフ＆ハーフ（塩／キャラメル）", 0.060, 550), ("ポップコーン", "北海道濃厚バターしょうゆ味", 0.030, 550),
    ("ポップコーン", "ポップコーン　塩Ｌ", 0.012, 650), ("ポップコーン", "ポップコーン　キャラメルＬ", 0.014, 650),
    ("その他フード", "アイスクレープ　バナナ＆チョコ", 0.004, 450),
    ("コールド", "ドリンクバー　コールド", 0.180, 475), ("コールド", "ミニッツメイドオレンジ", 0.010, 400),
    ("コーヒー", "ＣＯＳＴＡ　ホットコーヒー", 0.020, 450), ("コーヒー", "ＣＯＳＴＡ　アイスラテ", 0.015, 500),
    ("ホット", "ドリンクバー　ホット", 0.012, 475),
    ("アルコール", "プレミアムモルツ", 0.015, 700), ("アルコール", "ハイボール", 0.008, 650),
    ("その他ドリンク", "ＩＣＥＥ　イエロー＆ブルー", 0.020, 450), ("その他ドリンク", "超ストロベリーフローズン", 0.010, 550),
    ("ドリンク調味料", "ＥＣＯガムシロップ", 0.002, 0), ("フード調味料", "ケチャップソース３００Ｇ", 0.001, 0),
    ("ＳＥＴ作品コンボ", "サンプル作品　Ｄカップ　仕入", 0.008, 1200), ("ＳＥＴ作品コンボ", "サンプル作品　Ｐボックス　仕入", 0.005, 1500),
    ("引換券", "引換／ポップコーン塩Ｍ", 0.006, 0), ("引換券", "引換／ドリンクバー　アイス", 0.006, 0),
    ("コンセ包材", "Ｍカップ（ポップコーン）", 0.100, 0), ("コンセ包材", "バイオマスストロー", 0.180, 0),
]
CAT_CODE = {"コールド": "01", "コーヒー": "02", "ホット": "03", "アルコール": "04", "その他ドリンク": "05",
            "ドリンク調味料": "06", "ポップコーン": "31", "ホットドッグ": "32", "軽食系フード": "33",
            "調理系スイーツ": "34", "その他フード": "35", "フード調味料": "36", "ＳＥＴ作品コンボ": "51",
            "引換券": "61", "コンセ包材": "71"}
# MSO(注文明細)の店舗商品以外の商品: (商品名, 1日の明細数, 単価, 商品区分)。セット親は集計対象外
FILLER_MSO = [
    ("ドリンクバー　コールド", 330, 475, "単品"), ("ポップコーン　塩Ｍ", 70, 500, "単品"),
    ("ポップコーン　キャラメルＭ", 90, 500, "単品"), ("ハーフ＆ハーフ（塩／キャラメル）", 100, 550, "単品"),
    ("北海道濃厚バターしょうゆ味", 50, 550, "単品"), ("ＣＯＳＴＡ　ホットコーヒー", 35, 450, "単品"),
    ("ＩＣＥＥ　イエロー＆ブルー", 35, 450, "単品"), ("プレミアムモルツ", 25, 700, "単品"),
]
SET_PARENT = "ポップコーンセット"
SET_ITEMS = ["ハーフ＆ハーフ（塩／キャラメル）", "ポップコーン　塩Ｍ", "ポップコーン　キャラメルＭ"]
SET_DRINK = "ドリンクバー　コールド"
SET_ORDERS = 90                                 # 1日のセット注文数(セット親1行＋構成品2行)
# 時間帯ごとの注文の山(実データに近い形)。営業時間外の時間は係数算出の区切りに合わせて自動で外す
HOUR_W = {6: 20, 7: 40, 8: 300, 9: 110, 10: 190, 11: 900, 12: 180, 13: 510, 14: 950, 15: 400, 16: 520,
          17: 760, 18: 850, 19: 540, 20: 1070, 21: 530, 22: 190, 23: 240, 24: 60, 25: 20, 26: 15, 27: 10}


def product_code(name):
    """商品コード(13桁・仮)。売上CSVとMSO CSVで同じ商品は同じコードにする"""
    return f"18130009{zlib.crc32(name.encode('utf-8')) % 100000:05d}"


def rule_for(name):
    for r in PRODUCT_RULES:
        if r[0] and r[0] in name:
            return r
    return DEFAULT_RULE


def latest_sunday(today):
    return today - dt.timedelta(days=(today.weekday() - 6) % 7)


def serial_time(h, m):
    return (h * 60 + m) / 1440


def category_of(name):
    if "チュリトス" in name:
        return "調理系スイーツ"
    if "じゃが" in name or "チキン" in name:
        return "軽食系フード"
    if "ポップコーン" in name or "ハーフ＆ハーフ" in name or "しょうゆ味" in name or "バター味" in name:
        return "ポップコーン"
    return "ホットドッグ"


def store_products(wb):
    """期間データ!B14:B33 の登録商品(スロット順)。途中に空き枠があると verify の添字がずれるので止める"""
    pd = wb["期間データ"]
    out = []
    for i in range(N_SLOTS):
        v = pd[f"B{ROW_P0 + i}"].value
        out.append(v if isinstance(v, str) and v.strip() else None)
    filled = [i for i, v in enumerate(out) if v]
    if not filled:
        raise SystemExit("期間データ!B14:B33 に商品が登録されていません（店舗版ver2.0のファイルを指定してください）")
    if filled != list(range(len(filled))):
        raise SystemExit("期間データの商品枠に空きがあります（B14から詰めて登録してください）")
    return out[:len(filled)]


def hours_of(ws):
    """係数算出!C4:H4 の区切りを時(0〜48)で返す: (開店, 朝→昼, 昼→夕, 夕→夜, 夜→レイト, 閉店)。
    閉店が開店以前の時刻(翌2:00など)は +24h。シート側のN4/O4と同じ扱い"""
    vals = []
    for c in "CDEFGH":
        v = ws[f"{c}4"].value
        if isinstance(v, dt.datetime):
            vals.append(((v - dt.datetime(1899, 12, 30)).total_seconds() / 3600) if v.year >= 1900 else
                        v.hour + v.minute / 60)
        elif isinstance(v, dt.time):
            vals.append(v.hour + v.minute / 60 + v.second / 3600)
        elif isinstance(v, (int, float)):
            vals.append(float(v) * 24)
        else:
            raise SystemExit(f"係数算出!{c}4 が時刻ではありません（{v!r}）。店舗版の時間の区切りを確認してください")
    if vals[5] < 24 and vals[5] <= vals[0]:
        vals[5] += 24                                 # 翌日閉店
    if vals[5] >= 24 and vals[5] - 24 > vals[0]:
        vals[5] = 24 + vals[0]                        # シートは翌日分を開店でクランプ
    if not (vals[0] < vals[1] < vals[2] < vals[3] < vals[4] < vals[5]):
        raise SystemExit(f"係数算出!C4:H4 の区切りが順番になっていません: {vals}")
    return vals


# ---------------------------------------------------------------- 売上CSV ----
def make_sales_csv(path, products, start, end, att_sum, rng):
    """1商品1行(期間合計)の「売上・在庫・原価」CSVを書く。cp932・ヘッダー付き"""
    store = set(products)
    items = [(cat, name, rate, price) for cat, name, rate, price in FILLER_SALES if name not in store]
    for name in products:
        r = rule_for(name)
        items.append((category_of(name), name, r[1], r[5]))
    items.sort(key=lambda x: (CAT_CODE.get(x[0], "99"), x[1]))
    names = [x[1] for x in items]
    assert len(names) == len(set(names)), "売上CSVに同名商品が重複しています"
    rows = []
    for cat, name, rate, price in items:
        qty = max(0, int(round(rate * att_sum * rng.lognormvariate(0, 0.08))))
        stock0 = int(qty * 0.3) + rng.randint(0, 5)
        purchase = qty + rng.randint(0, 3)
        waste = rng.randint(0, 2) if price else 0
        left = stock0 + purchase - qty - waste
        gross = qty * price
        net = round(gross / 1.08) if price else 0
        rows.append(["売上・在庫・原価", THEATER[0], THEATER[1], start.strftime("%Y%m%d"), end.strftime("%Y%m%d"),
                     "01", CAT_CODE.get(cat, "99"), cat, "", "",
                     "19900001", "サンプル商事株式会社", product_code(name), name, "1", "買取",
                     stock0, purchase, 0, 0, waste, 0, 0, 0, 0, left, qty, gross, qty if price else 0,
                     gross, net, round(price / 1.08, 2) if price else 0, round(net * 0.35), "0"])
    with open(path, "w", encoding="cp932", newline="") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADERS)
        w.writerows(rows)
    return {r[13]: r[26] for r in rows}


# ---------------------------------------------------------------- MSO CSV ----
def make_mso_csv(path, products, day, hours, rng):
    """金曜1日分のMSO商品CSV(注文明細)を書く。時刻の分布は実データに近い山を持たせ、
    店舗商品は商品ごとに帯別の重みを変えて「商品別の波」が見えるようにする。
    時刻は営業日ベースの秒(翌0〜2時は24時以降)で扱い、伝票番号は時系列に振る"""
    open_h, *cuts, close_h = hours
    store = set(products)

    def band_index(h):
        mid = h + 0.5
        for bi, c in enumerate(cuts):
            if mid < c:
                return bi
        return 4

    allowed = []                                  # (時, 分の下限, 分の上限)
    for h in range(int(open_h), int(close_h - 1e-9) + 1):
        lo = max(0, int(round((open_h - h) * 60))) if h == int(open_h) else 0
        hi = min(59, int(round((close_h - h) * 60)) - 1) if h == int(close_h - 1e-9) and close_h < h + 1 else 59
        if lo <= hi:
            allowed.append((h, lo, hi))

    def sample_time(mult=(1, 1, 1, 1, 1)):
        w = [HOUR_W.get(h, 20) * mult[band_index(h)] for h, _, _ in allowed]
        h, lo, hi = rng.choices(allowed, w)[0]
        return h * 3600 + rng.randint(lo, hi) * 60 + rng.randint(0, 59)

    lines = []                                   # (営業日秒, 商品名, 単価, 数量, 区分, セットID)
    for name in products:
        r = rule_for(name)
        n = max(0, int(round(r[3] * rng.lognormvariate(0, 0.12))))
        for _ in range(n):
            lines.append([sample_time(r[4]), name, r[5], 1 if rng.random() < 0.85 else 2, "単品", None])
    for name, n_day, price, kind in FILLER_MSO:
        if name in store:
            continue                              # 店舗商品と同名の埋め草は出さない(二重計上防止)
        n = max(0, int(round(n_day * rng.lognormvariate(0, 0.1))))
        for _ in range(n):
            lines.append([sample_time(), name, price, 1 if rng.random() < 0.8 else 2, kind, None])
    set_items = [n for n in SET_ITEMS if n not in store] or SET_ITEMS
    for k in range(SET_ORDERS):                  # セット親＋構成品(親は集計対象外、構成品は対象)
        t = sample_time()
        gid = f"set{k}"
        lines.append([t, SET_PARENT, 950, 1, "セット親", gid])
        lines.append([t, rng.choice(set_items), 475, 1, "構成品", gid])
        lines.append([t, SET_DRINK, 475, 1, "構成品", gid])
    lines.sort(key=lambda x: (x[0], x[5] or ""))
    rows = []
    slip = 0
    i = 0
    ymd = day.strftime("%Y/%m/%d")
    while i < len(lines):
        grp = [lines[i]]
        if lines[i][5]:                          # セット注文は3行で1伝票
            while i + len(grp) < len(lines) and lines[i + len(grp)][5] == lines[i][5]:
                grp.append(lines[i + len(grp)])
        else:
            while (len(grp) < 3 and i + len(grp) < len(lines) and not lines[i + len(grp)][5]
                   and rng.random() < 0.35):
                grp.append(lines[i + len(grp)])
        slip += 1
        mobile = rng.random() < 0.08
        t0 = grp[0][0]
        order_day = day + dt.timedelta(days=t0 // 86400)          # 翌0〜2時は翌日付の注文
        total = sum(p * q for _, _, p, q, kind, _ in grp if kind != "構成品")
        cancel = rng.random() < 0.004            # ごく少数の取消(集計対象外)
        for j, (t, name, price, qty, kind, _) in enumerate(grp, start=1):
            tt = [t, t + 25, t + 190, t + 200]
            rows.append([1, f"0764{day.strftime('%y%m%d')}{slip:05d}", "01", "000", "サンプル", THEATER[0], THEATER[1],
                         "0080" if mobile else "0070", "モバイルオーダー" if mobile else "セルフオーダー",
                         40 if mobile else rng.randint(71, 78), ymd, order_day.strftime("%Y/%m/%d"),
                         *(f"{(x // 3600) % 24:02d}:{(x // 60) % 60:02d}:{x % 60:02d}" for x in tt),
                         total, slip, "提供済", "取消" if cancel else "販売",
                         "0000000001" if mobile else "",
                         j, product_code(name), name, price, qty,
                         0 if kind == "構成品" else price * qty, kind])
        i += len(grp)
    with open(path, "w", encoding="cp932", newline="") as f:
        w = csv.writer(f)
        w.writerow(MSO_HEADERS)
        w.writerows(rows)
    return len(rows)


# ---------------------------------------------------------------- 貼り込み ----
def paste(ws, rows, ncol, max_rows):
    for i, row in enumerate(rows[:max_rows]):
        for c_idx, v in enumerate(row[:ncol], start=1):
            ws.cell(row=5 + i, column=c_idx).value = v


def build(src, dst, csv_dir, sunday, label=True, password=None):
    rng = random.Random(SEED)
    os.makedirs(csv_dir, exist_ok=True)
    a_end = sunday
    a_start = a_end - dt.timedelta(days=CSV_A_DAYS - 1)              # 金
    b_end = a_start - dt.timedelta(days=1)                            # 前週木
    b_start = b_end - dt.timedelta(days=CSV_B_DAYS - 1)               # 前週金
    fridays = [a_start - dt.timedelta(days=7 * k) for k in (3, 2, 1, 0)]   # 4週分(古い順)
    assert a_start.weekday() == 4 and b_start.weekday() == 4
    stale_from = b_end + dt.timedelta(days=STALE_DAYS + 1)            # 「古いデータ」通知が出始める日(期間B側が先)

    wb = load_workbook(src)
    names = store_products(wb)
    ks = wb["係数算出"]
    hours = hours_of(ks)
    csv_a = f"sample_periodA_{a_start:%m%d}-{a_end:%m%d}.csv"
    csv_b = f"sample_periodB_{b_start:%m%d}-{b_end:%m%d}.csv"
    make_sales_csv(os.path.join(csv_dir, csv_a), names, a_start, a_end, sum(ATT_A), rng)
    make_sales_csv(os.path.join(csv_dir, csv_b), names, b_start, b_end, sum(ATT_B), rng)
    mso_files = []
    for k, d in enumerate(fridays):
        fn = f"sample_mso_week{k + 1}_{d:%m%d}.csv"
        make_mso_csv(os.path.join(csv_dir, fn), names, d, hours, rng)
        mso_files.append(fn)

    paste(wb["CSV貼付A"], read_csv_rows(os.path.join(csv_dir, csv_a)), NCOL, CSV_MAX)
    paste(wb["CSV貼付B"], read_csv_rows(os.path.join(csv_dir, csv_b)), NCOL, CSV_MAX)
    for sheet, fn in zip(CALIB_SHEETS, mso_files):
        paste(wb[sheet], read_mso_rows(os.path.join(csv_dir, fn)), MSO_NCOL, MSO_MAX)

    pd = wb["期間データ"]
    for c, v in zip("CDE", ATT_A):
        pd[f"{c}6"] = v
    for c, v in zip("CDEFGHI", ATT_B):
        pd[f"{c}10"] = v
    pd["E12"] = HT_THR
    ht, manual = [], {}
    for i, name in enumerate(names):
        r = ROW_P0 + i
        rule = rule_for(name)
        pd[f"E{r}"] = rule[2]
        ht.append(rule[2])
        for key, v in MANUAL_PRIORITY.items():
            if key in name:
                pd[f"F{r}"] = v
                manual[i] = v
        for key, v in MEMOS.items():
            if key in name:
                pd[f"G{r}"] = v

    m = wb["準備数計算"]
    m["D4"] = "期間A（直近金土日）"
    m["D5"] = PEAK_ATT
    f4, g4 = ks["F4"].value, ks["G4"].value                       # ④ 夜ピークの区切り(店舗設定)
    if not all(isinstance(x, (dt.time, dt.datetime)) for x in (f4, g4)):
        raise SystemExit("係数算出!F4/G4 が時刻ではないため、時間帯プリセット名（④夜ピーク）を作れません")
    m["D6"] = f"④ 夜ピーク{f4.hour}:{f4.minute:02d}〜{g4.hour}:{g4.minute:02d}"   # 係数算出!P10 と同じ文字列
    m["D7"] = "使う"
    m["D8"] = RATE
    m["O5"] = serial_time(*PEAK_START)
    m["O6"] = serial_time(*PEAK_END)
    wb["印刷用"][VIEW_ADDR] = VIEW_ALL
    if label:                                    # 実データと取り違えないよう全シートに明記
        guide = (f"このサンプルは①〜③まで入力済みです（期間A {a_start.month}/{a_start.day}〜{a_end.month}/{a_end.day}・"
                 f"期間B {b_start.month}/{b_start.day}〜{b_end.month}/{b_end.day}・MSO金曜4週・動員数・保持時間・"
                 f"⑥ピーク時間 {PEAK_START[0]}:{PEAK_START[1]:02d}〜・⑤事前準備率 {RATE}% はすべて仮の値）。"
                 f"貼り替えの練習は同梱の sample_csv のCSVで。※日付は作成時のもので、{stale_from.month}/{stale_from.day} 以降は"
                 "「古いデータ」の注意が出ますが仮データなのでそのままで構いません。")
        targets = [("使い方", "B2"), ("準備数計算", "B2"), ("期間データ", "F2"),
                   ("印刷用", "B5"), ("係数算出", "B2"), ("商品別の波", "B2")]
        targets += [(s, "A2") for s in ("CSV貼付A", "CSV貼付B", *CALIB_SHEETS)]
        for sheet, ref in targets:
            ws = wb[sheet]
            old = str(ws[ref].value or "").lstrip("｜").strip()
            ws[ref] = TAG + "　" + old
        wb["印刷用"]["B1"] = str(wb["印刷用"]["B1"].value or "") + "【サンプル】"
        # 使い方の空き行(3行目)に「入力済み」の案内を折り返しで置く(B2に足すと印刷が縮むため)
        from openpyxl.styles import Alignment, Font
        ws = wb["使い方"]
        if not any(str(r) == "B3:J3" for r in ws.merged_cells.ranges):
            ws.merge_cells("B3:J3")
        ws["B3"] = guide
        ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
        ws["B3"].font = Font(size=9, color="7A4A00")
        ws.row_dimensions[3].height = 42
    lock_workbook(wb, password)                  # 元ファイルと同じ保護状態で保存する
    wb.save(dst)
    n = restore_comment_vml(src, dst)
    ok, hidden = check_hidden_cols(dst)
    if not ok:
        raise SystemExit(f"非表示列が壊れています: {hidden}")

    params = {
        "csv_a": csv_a, "csv_b": csv_b, "mso": mso_files, "att_a": ATT_A, "att_b": ATT_B,
        "peak": PEAK_ATT, "rate": RATE, "ht": ht, "manual": manual, "thr": HT_THR,
        "peak_start": f"{PEAK_START[0]}:{PEAK_START[1]:02d}", "band": 4, "products": names,
        "period_a": [a_start.isoformat(), a_end.isoformat()], "period_b": [b_start.isoformat(), b_end.isoformat()],
        "fridays": [d.isoformat() for d in fridays], "stale_notice_from": stale_from.isoformat(),
        "note": "CSVのパスはこのファイルのあるフォルダからの相対パス",
    }
    with open(os.path.join(csv_dir, "sample_params.json"), "w", encoding="utf-8") as f:
        json.dump(params, f, ensure_ascii=False, indent=1)
    args = ["--products", "auto", "--csv-a", csv_a, "--csv-b", csv_b,
            "--att-a", ",".join(map(str, ATT_A)), "--att-b", ",".join(map(str, ATT_B)),
            "--peak", str(PEAK_ATT), "--rate", str(RATE), "--view", VIEW_ALL,
            "--ht", ",".join(str(h) for h in ht), "--thr", str(HT_THR), "--peak-start", params["peak_start"]]
    if manual:
        args += ["--manual", ",".join(f"{k}:{v}" for k, v in manual.items())]
    for k, fn in enumerate(mso_files):
        args += [f"--mso{k + 1}", fn]
    # $(cat verify_args.txt) でそのまま展開できるよう空白を含まないことを確認して素の連結で書く
    # (shlex.join は日本語を引用符で囲み、$(cat) 経由では引用符が残るため使わない)
    bad = [x for x in args if any(ch.isspace() for ch in x)]
    if bad:
        raise SystemExit(f"検証引数に空白を含む値があります: {bad}")
    with open(os.path.join(csv_dir, "verify_args.txt"), "w", encoding="utf-8") as f:
        f.write(" ".join(args) + "\n")
    print(f"sample: {dst}  (products={len(names)}, 期間A {a_start:%m/%d}〜{a_end:%m/%d}, "
          f"期間B {b_start:%m/%d}〜{b_end:%m/%d}, MSO {', '.join(f'{d:%m/%d}' for d in fridays)}, "
          f"事前準備率 {RATE}%, 「古いデータ」通知は {stale_from:%Y/%m/%d} から, comments restored: {n})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--csv-dir", default=None, help="合成CSVと検証引数の出力先(既定: 出力xlsxと同じフォルダの sample_csv)")
    ap.add_argument("--sunday", default=None, help="期間Aの最終日(日曜)。省略時は直近の日曜")
    ap.add_argument("--no-label", action="store_true", help="「サンプル」の表記を入れない")
    ap.add_argument("--password", help="シート保護の解除パスワード")
    a = ap.parse_args()
    sun = dt.date.fromisoformat(a.sunday) if a.sunday else latest_sunday(dt.date.today())
    if sun.weekday() != 6:
        raise SystemExit("--sunday は日曜の日付を指定してください")
    build(a.src, a.dst, a.csv_dir or os.path.join(os.path.dirname(os.path.abspath(a.dst)), "sample_csv"), sun,
          label=not a.no_label, password=a.password)
