# -*- coding: utf-8 -*-
"""
v2.0(upgrade_v2.py で更新した店舗版)の再計算結果を、CSVからPythonで独立計算した
期待値と突合する。事前準備率・保持時間/作るタイミング・印刷用の表示切替を検証する。
  python verify_v2.py <再計算済みxlsx> --csv-a A.csv --csv-b B.csv --att-a 9000,13000,12000 \
      --att-b ... --peak 1200 --mult 0.6 --rate 80 --view 直前に作るもの \
      --ht 10,45,,30,... --manual 1:直前,5:先に --thr 30 [--expect-warn 文字列] [--mso]
  --products auto : 商品名を期間データ!B14:B33 から読む(店舗版・サンプル用。既定はテンプレートの商品。
                    枠はB14から詰めて登録されている前提。途中に空きがあるとNG)
  --mso1〜4 CSV   : 係数貼付①〜④に貼ったMSO商品CSV。係数算出・商品別の波・係数の適用まで厳密に突合
                    (--mso は貼付ありの近似検証のみ)。指定時は --mult 不要(実測係数を期待値にする)
  --manual        : 手動のタイミング。添字は0始まり(0＝期間データB14の1商品目)。例 --manual 0:直前,4:先に
                    旧表記の 高/低 を渡すと 直前/先に に読み替えて期待値にする(後方互換の確認用)
  --peak-start    : ⑥ピーク開始(準備数計算!O5)。既定 17:30。O5が未入力のブックは --peak-start "" (仕込み開始=—)
  ※ 商品名の突合はPython側は完全一致。シートのSUMIFSは大文字小文字を区別せず *?~ をワイルドカードに
    するため、そのような商品名は結果が食い違う(NGとして出る)
"""
import argparse
import math
import sys

import re as _re

from openpyxl import load_workbook

from build_calib import CALIB_SHEETS, MSO_MAX, read_mso_rows
from build_tool import (DEFAULT_PRODUCTS, JUDGE_FEW, JUDGE_NODATA, JUDGE_NONE, JUDGE_USE, N_SLOTS,
                        ROW_M0, WAVE_ROW0, WAVE_SHEET, read_csv_rows)
from upgrade_v2 import (PRI_HI, PRI_LO, PRI_OLD, RATE, VIEW_CELL, VIEW_HI, VIEW_LO,  # noqa: E402
                        ftext)

ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--csv-a", required=True)
ap.add_argument("--csv-b", required=True)
ap.add_argument("--att-a", required=True)
ap.add_argument("--att-b", required=True)
ap.add_argument("--peak", type=int, required=True)
ap.add_argument("--mult", type=float, default=None, help="期待する時間帯係数(M7)。--mso1〜4指定時は不要(実測係数を期待)")
ap.add_argument("--rate", default="100", help="事前準備率(整数%) か 'blank'(=100扱い+⚠)")
ap.add_argument("--view", default="すべて")
ap.add_argument("--ht", required=True, help="保持時間(商品順・空欄は空文字)")
ap.add_argument("--manual", default="",
                help="手動のタイミング: 添字(0始まり・0＝期間データB14の1商品目):直前/先に をカンマ区切り(旧 高/低 も可)")
ap.add_argument("--thr", type=int, default=30)
ap.add_argument("--expect-warn", action="append", default=[])
ap.add_argument("--mso", action="store_true", help="MSO貼付あり(商品係数で販売予測数が変わるため近似検証。厳密は --mso1〜4)")
for _k in range(1, 5):
    ap.add_argument(f"--mso{_k}", help=f"係数貼付{'①②③④'[_k - 1]}に貼ったMSO商品CSV(厳密検証)")
ap.add_argument("--peak-start", default="17:30", help="⑥ピーク開始(空文字=未入力)")
ap.add_argument("--products", default="default", choices=["default", "auto"])
a = ap.parse_args()

errors = []


def check(label, got, want, tol=0):
    ok = (got == want) if tol == 0 else (
        isinstance(got, (int, float)) and isinstance(want, (int, float)) and abs(got - want) <= tol)
    if not ok:
        errors.append(f"NG {label}: got={got!r} want={want!r}")


def fl(x):
    return max(0, math.floor(x + 1e-9))


def as_serial(v):
    """時刻書式のセルは datetime.time/datetime で読めるので日のシリアル値へ"""
    import datetime as _d
    if isinstance(v, _d.datetime):
        # openpyxlは 0<シリアル<60 を1899/12/31起点で変換する(1900年閏年バグの扱い)。負や60以上は1899/12/30起点
        epoch = (_d.datetime(1899, 12, 31) if _d.datetime(1900, 1, 1) <= v < _d.datetime(1900, 3, 1)
                 else _d.datetime(1899, 12, 30))
        return (v - epoch).total_seconds() / 86400
    if isinstance(v, _d.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v


def check_r005(label, got, raw):
    """0.05刻みセルの検証。刻みのちょうど中間は浮動小数の微差で丸め方向が
    環境依存になるため、「0.05の倍数」かつ「生係数から半ステップ以内」を確認する"""
    ok = (isinstance(got, (int, float))
          and abs(got / 0.05 - round(got / 0.05)) < 1e-6
          and abs(got - raw) <= 0.025 + 1e-6)
    if not ok:
        errors.append(f"NG {label}: got={got!r} raw={raw!r}")


wb = load_workbook(a.xlsx, data_only=True)
wf = load_workbook(a.xlsx)
m, pr, pd = wb["準備数計算"], wb["印刷用"], wb["期間データ"]
mf = wf["準備数計算"]

if a.products == "auto":
    slots = [wf["期間データ"][f"B{14 + i}"].value for i in range(N_SLOTS)]
    slots = [v if isinstance(v, str) and v.strip() else None for v in slots]
    PRODUCTS = [v for v in slots if v]
    if slots[:len(PRODUCTS)] != PRODUCTS:
        gap = next(i for i, v in enumerate(slots) if v is None)
        print(f"FAILED (1):\n  NG 期間データの商品枠に空きがあります（B{14 + gap} が空欄で、その後に商品があります）")
        sys.exit(1)
    if not PRODUCTS:
        print("FAILED (1):\n  NG 期間データ!B14:B33 に商品がありません")
        sys.exit(1)
else:
    PRODUCTS = list(DEFAULT_PRODUCTS)
NP = len(PRODUCTS)

sales = {}
for r in read_csv_rows(a.csv_a):
    if r[13]:
        sales[r[13]] = sales.get(r[13], 0) + (r[26] or 0)
att = sum(int(x) for x in a.att_a.split(","))
ht = [None if s == "" else int(s) for s in a.ht.split(",")]
ht += [None] * (NP - len(ht))
_pri_norm = {PRI_HI: PRI_HI, PRI_LO: PRI_LO, PRI_OLD[PRI_HI]: PRI_HI, PRI_OLD[PRI_LO]: PRI_LO}
manual = {int(k): _pri_norm[v] for k, v in (kv.split(":") for kv in a.manual.split(",") if kv)}
rate = 100 if a.rate == "blank" or int(a.rate) <= 0 else int(a.rate)   # 0以下は100%扱い(⚠つき)
peak = None
if a.peak_start:
    _ph, _pm = a.peak_start.split(":")
    peak = (int(_ph) * 60 + int(_pm)) / 1440

# ---- 数式レベル: G列はF列の式に率を掛けただけ(端数処理・係数条件は同一)、参照の付け替え
for i in range(N_SLOTS):
    r = ROW_M0 + i
    f_txt, g_txt = ftext(mf[f"F{r}"]), ftext(mf[f"G{r}"])
    rnd = _re.search(r"ROUND(?:UP|DOWN)?", f_txt).group(0)
    want_g = (f'=IF(OR($C{r}="",NOT(ISNUMBER($D$5))),"",'
              f'IF(ISNUMBER($F{r}),MAX(0,{rnd}($F{r}*{RATE}/100,0)),"—"))')
    check(f"G{r}式=表示中の販売予測数×率", g_txt, want_g)
    check(f"P{r}式(手動はTRIM判定)", "TRIM(期間データ!F" in ftext(mf[f"P{r}"]), True)
b9f = ftext(mf["B9"])
check("B9式に旧G4参照なし", "期間データ!$G$4" in b9f, False)
check("B9式にB7参照", "期間データ!$B$7" in b9f, True)
check("B9式に比較期間はN列案内", "G列「比較期間" in b9f, False)
check("印刷用!B2式に率", RATE.replace("$D$8", "準備数計算!$D$8") in ftext(wf["印刷用"]["B2"]), True)
check("L列式が印刷用の切替セルを参照", VIEW_CELL in ftext(mf[f"L{ROW_M0}"]), True)
check("Q列式(開始目安=MOD)", "MOD($K" in ftext(mf[f"Q{ROW_M0}"]), True)
check("印刷用!B4式にピーク表示", "ピーク" in ftext(wf["印刷用"]["B4"]), True)
pp = wf["印刷用"]
check("印刷用 A4縦1枚(fitToPage)", bool(pp.sheet_properties.pageSetUpPr and pp.sheet_properties.pageSetUpPr.fitToPage), True)
check("印刷用 fitToWidth/Height=1", (pp.page_setup.fitToWidth, pp.page_setup.fitToHeight), (1, 1))
check("印刷用 A4/縦", (pp.page_setup.paperSize, pp.page_setup.orientation), (9, "portrait"))
check("印刷用 印刷範囲", pp.print_area, "'印刷用'!$A$1:$I$28")
check("印刷用!E列(調理の目安)がINDIRECTで調理時間を参照", "調理時間" in ftext(wf["印刷用"]["E8"]), True)

# ---- MSO(係数算出・商品別の波)の厳密検証: 帯の区切りは係数算出シートの設定から読む
mso_paths = [getattr(a, f"mso{k}") for k in range(1, 5)]
exact_mso = any(mso_paths)
if not exact_mso and a.mult is None:
    ap.error("--mult は --mso1〜4 を指定しないときに必要です")
eff = {}                       # 商品名 → 販売予測数に掛かる係数(商品係数 or 全体係数)
if exact_mso:
    ks, wv = wb["係数算出"], wb[WAVE_SHEET]
    c4, d4, e4, f4, g4, h4 = (as_serial(ks[f"{c}4"].value) for c in "CDEFGH")
    if h4 >= 1:
        n4, o4 = min(h4 % 1, c4), 1
    elif h4 >= g4:
        n4, o4 = 0, h4
    elif h4 <= c4:
        n4, o4 = h4, 1
    else:
        n4, o4 = 0, h4
    check("係数算出!N4", as_serial(ks["N4"].value), n4, tol=1e-9)
    check("係数算出!O4", as_serial(ks["O4"].value), o4, tol=1e-9)
    DURS = [(d4 - c4) * 24, (e4 - d4) * 24, (f4 - e4) * 24, (g4 - f4) * 24, (o4 - g4) * 24 + n4 * 24]
    windows = [(c4, d4), (d4, e4), (e4, f4), (f4, g4)]
    if any(d <= 0 for d in DURS) or sum(DURS) <= 0:
        # シート側は帯長0以下を"—"/⚠で扱う。ここでは厳密検証の前提が崩れているのでNGにして止める
        print(f"FAILED (1):\n  NG 係数算出!C4:H4 の時間の区切りが不正です(帯の長さ {DURS})。厳密検証は区切りが正しいブックのみ")
        sys.exit(1)

    def band_of(tv):
        for bi, w in enumerate(windows):
            if w[0] <= tv < w[1]:
                return bi
        if (g4 <= tv < o4) or tv < n4:
            return 4
        return None

    def mso_rows_filtered(path):
        """Excel側AF列と同じ対象行(先頭行の日付・提供済・販売・セット親以外・数量>0)"""
        rows = read_mso_rows(path)[:MSO_MAX]
        target = rows[0][10] if rows else None
        out = []
        for r in rows:
            if r[10] != target or r[18] != "提供済" or r[19] != "販売" or r[27] == "セット親":
                continue
            q = r[25] or 0
            if not isinstance(q, (int, float)) or q <= 0:
                continue
            t = r[12]
            if isinstance(t, (int, float)):
                tv = t % 1.0
            elif isinstance(t, str):
                try:
                    parts = [int(x) for x in t.split(":")]
                except ValueError:
                    continue
                if len(parts) == 2:
                    parts.append(0)
                if len(parts) != 3:
                    continue
                tv = ((parts[0] * 3600 + parts[1] * 60 + parts[2]) / 86400) % 1.0
            else:
                continue
            out.append((tv, q, r[23]))
        return out

    week_counts, prod_counts = [], {}
    for p in mso_paths:
        wcs = [0] * 5
        if p:
            for tv, q, name in mso_rows_filtered(p):
                b = band_of(tv)
                if b is None:
                    continue
                wcs[b] += q
                if name:
                    prod_counts.setdefault(name, [0] * 5)[b] += q
        week_counts.append(wcs)
    n_weeks = sum(1 for wcs in week_counts if sum(wcs) > 0)
    wave_thr = wv["D3"].value if isinstance(wv["D3"].value, (int, float)) else 30

    for wi, wcs in enumerate(week_counts):
        col = "DEFG"[wi]
        for bi in range(5):
            check(f"係数算出!{col}{7 + bi}({wi + 1}週目)", ks[f"{col}{7 + bi}"].value, wcs[bi])
        check(f"係数算出!{col}12(週合計)", ks[f"{col}12"].value, sum(wcs))
        st = wb[CALIB_SHEETS[wi]]["A3"].value or ""
        if mso_paths[wi]:
            check(f"{CALIB_SHEETS[wi]}!AD3(正常フラグ)", wb[CALIB_SHEETS[wi]]["AD3"].value, 1)
            check(f"{CALIB_SHEETS[wi]}!A3(貼付済表示)", "貼付" in st and "対象個数" in st, True)
            check(f"{CALIB_SHEETS[wi]}!A3(金曜)", "金曜ではありません" in st, False)
        else:
            check(f"{CALIB_SHEETS[wi]}!A3(未貼付表示)", "未貼付" in st, True)
    check("係数算出!H13(週数)", ks["H13"].value, n_weeks)
    band_coef = [None] * 5
    # 店舗版はJ列(係数)が =IF(ISNUMBER(K),K,既定) で係数算出の実測を自動採用する。テンプレート系は
    # J列が定数(手で転記)なので、J式がK列を参照するときだけ「J=実測」を要求する
    j_linked = all("K" in ftext(mf[f"J{4 + bi}"]) for bi in range(5))
    if n_weeks:
        avgs = [sum(wcs[bi] for wcs in week_counts) / n_weeks for bi in range(5)]
        day_pace = sum(avgs) / sum(DURS)
        for bi in range(5):
            coef = (avgs[bi] / DURS[bi]) / day_pace
            check(f"係数算出!K{7 + bi}(係数候補)", ks[f"K{7 + bi}"].value, coef, tol=1e-6)
            lv = ks[f"L{7 + bi}"].value
            check_r005(f"係数算出!L{7 + bi}(転記用)", lv, coef)
            check(f"準備数計算!K{4 + bi}(実測候補の連動)", m[f"K{4 + bi}"].value, lv, tol=1e-9)
            if j_linked:
                check(f"準備数計算!J{4 + bi}(係数=実測)", m[f"J{4 + bi}"].value, lv, tol=1e-9)
            band_coef[bi] = lv
        check("係数算出!B23(警告なし)", ks["B23"].value in (None, ""), True)
        check("係数算出!N5(区切り外=0)", ks["N5"].value, 0)
        check("係数算出!O5(同日疑い週=0)", ks["O5"].value, 0)
    sel = m["M8"].value
    check("準備数計算!M8(選択帯1〜5)", isinstance(sel, (int, float)) and 1 <= sel <= 5, True)
    sel = int(sel) if isinstance(sel, (int, float)) else 1
    d6 = str(m["D6"].value or "").strip()
    check("準備数計算!M8(D6の先頭マークと一致)", sel, "①②③④⑤".find(d6[:1]) + 1 if d6 else 0)
    if n_weeks and j_linked:
        mult = band_coef[sel - 1]
    else:                                        # 実測なし or J列が定数: シートのJ(選択帯)がそのまま係数
        jv = m[f"J{3 + sel}"].value
        check(f"準備数計算!J{3 + sel}(選択帯の係数が数値)", isinstance(jv, (int, float)), True)
        mult = jv if isinstance(jv, (int, float)) else (a.mult or 1)
    check("準備数計算!M7(選択帯の係数)", m["M7"].value, mult, tol=1e-9)
    check("準備数計算!M9(波の適用)", m["M9"].value, 1 if n_weeks else 0)
    for i, name in enumerate(PRODUCTS):
        wr = WAVE_ROW0 + i
        cs = prod_counts.get(name, [0] * 5)
        total = sum(cs)
        check(f"{WAVE_SHEET}!B{wr}(商品名)", wv[f"B{wr}"].value, name)
        for bi in range(5):
            check(f"{WAVE_SHEET}!{'PQRST'[bi]}{wr}({name[:5]})", wv[f"{'PQRST'[bi]}{wr}"].value, cs[bi])
            ccol = "CDEFG"[bi]
            if total == 0:
                check(f"{WAVE_SHEET}!{ccol}{wr}(構成比—)", wv[f"{ccol}{wr}"].value, "—")
            else:
                check(f"{WAVE_SHEET}!{ccol}{wr}(構成比)", wv[f"{ccol}{wr}"].value, cs[bi] / total, tol=1e-9)
        check(f"{WAVE_SHEET}!H{wr}(合計)", wv[f"H{wr}"].value, total)
        want_j = (JUDGE_NODATA if n_weeks == 0 else JUDGE_NONE if total == 0 else
                  JUDGE_FEW if total < wave_thr else JUDGE_USE)
        check(f"{WAVE_SHEET}!N{wr}(判定)", wv[f"N{wr}"].value, want_j)
        use = None
        for bi in range(5):
            kcol = "IJKLM"[bi]
            raw = None
            if want_j == JUDGE_USE and cs[bi] > 0 and DURS[bi] > 0:
                raw = (cs[bi] / total) * (sum(DURS) / DURS[bi])
                if raw / 0.05 < 0.5 - 1e-9:
                    raw = None
            got = wv[f"{kcol}{wr}"].value
            if raw is None:
                check(f"{WAVE_SHEET}!{kcol}{wr}(係数—)", got, "—")
            else:
                check_r005(f"{WAVE_SHEET}!{kcol}{wr}(係数)", got, raw)
                if bi == sel - 1 and isinstance(got, (int, float)):
                    use = got
        eff[name] = use if use is not None else mult
        check(f"準備数計算!H{ROW_M0 + i}(商品係数表示)", m[f"H{ROW_M0 + i}"].value, use if use is not None else "—",
              tol=1e-9 if use is not None else 0)
elif a.mso:
    check("準備数計算!M7(実測係数が数値)", isinstance(m["M7"].value, (int, float)) and m["M7"].value > 0, True)
    check("準備数計算!M9(波の適用)", m["M9"].value, 1)
else:
    check("準備数計算!M7", m["M7"].value, a.mult, tol=1e-9)
check("準備数計算!M4(期間A)", m["M4"].value, 1)
check("期間データ!E12(基準)", pd["E12"].value, a.thr)
check("E8に率表示", f"× {rate}%" in (m["E8"].value or ""), True)
check("印刷用!B2に率表示", f"{rate}%" in (pr["B2"].value or ""), True)

pri, starts = [], []
for i, name in enumerate(PRODUCTS):
    r = ROW_M0 + i
    rt = sales.get(name, 0) / att
    check(f"C{r}(商品名)", m[f"C{r}"].value, name)
    check(f"D{r}(期間販売数)", m[f"D{r}"].value, sales.get(name, 0))
    f_got, g_got = m[f"F{r}"].value, m[f"G{r}"].value
    if exact_mso:
        want_f = fl(a.peak * rt * eff[name])
        check(f"F{r}(販売予測数・係数{eff[name]:.2f})", f_got, want_f, tol=1)
        check(f"G{r}(作る数=販売予測数×率)", g_got,
              fl(f_got * rate / 100) if isinstance(f_got, (int, float)) else want_f)
    elif a.mso:
        # 商品係数は商品別の波(別途検証済み)に依存するため、率の掛かり方だけ確認
        check(f"F{r}数値", isinstance(f_got, (int, float)), True)
        if isinstance(f_got, (int, float)):
            check(f"G{r}=F×率", g_got, fl(f_got * rate / 100))
    else:
        check(f"F{r}(販売予測数)", f_got, fl(a.peak * rt * a.mult), tol=1)
        check(f"G{r}(作る数=販売予測数×率)", g_got,
              fl(f_got * rate / 100) if isinstance(f_got, (int, float)) else fl(a.peak * rt * a.mult * rate / 100))
    h = ht[i]
    check(f"O{r}(保持時間)", m[f"O{r}"].value, "—" if h is None else h)
    want_p = manual.get(i) or ("—" if h is None else (PRI_HI if h <= a.thr else PRI_LO))
    check(f"P{r}(作るタイミング)", m[f"P{r}"].value, want_p)
    pri.append(want_p)
    # 仕込み開始の目安(Q)= MOD(ピーク開始 − 保持時間, 1)。K(非表示)は未補正
    if peak is not None and h is not None:
        start = peak - h / 1440
        check(f"Q{r}(開始目安)", as_serial(m[f"Q{r}"].value), start % 1, tol=1e-6)
        check(f"K{r}(開始・未補正)", as_serial(m[f"K{r}"].value), start, tol=1e-6)
        starts.append(round(start * 1440))
    else:
        check(f"Q{r}(開始目安=—)", m[f"Q{r}"].value, "—")
        starts.append(None)
    # 並び順キー(非表示L列): 保持時間の長い順、未入力は最後(ピーク入力の有無によらない)
    k = i + 1
    hkey = (9999900 + k) if h is None else (10000 - max(0, min(9999, h))) * 100 + k
    if a.view == VIEW_HI:
        want_l = "" if want_p == PRI_LO else hkey
    elif a.view == VIEW_LO:
        want_l = hkey if want_p == PRI_LO else ""
    else:
        want_l = hkey
    got_l = m[f"L{r}"].value
    check(f"L{r}(並び順キー)", "" if got_l in (None, "") else got_l, want_l)
for i in range(NP, N_SLOTS):                   # 未登録の枠は空
    r = ROW_M0 + i
    for c in "CDEFGHNOPQ":
        check(f"{c}{r}(空枠)", m[f"{c}{r}"].value in (None, ""), True)
    check(f"L{r}(空枠キー)", m[f"L{r}"].value in (None, ""), True)
    if exact_mso:
        wr = WAVE_ROW0 + i
        for c in "BCDEFGHIJKLMNPQRST":
            check(f"{WAVE_SHEET}!{c}{wr}(空枠)", wb[WAVE_SHEET][f"{c}{wr}"].value in (None, ""), True)

# 印刷用「調理の目安」= 調理時間シートの所要時間(分)を切り上げ。シートが無い/該当なしは「—」
cook_min = {}
if "調理時間" in wb.sheetnames:
    import add_cooking_sheet as _CK
    _ck = wb["調理時間"]
    for _r in range(_CK.ROW_C0, _CK.ROW_C0 + _CK.N_ROWS):
        _n, _v = _ck[f"{_CK.COL_NAME}{_r}"].value, _ck[f"{_CK.COL_MIN}{_r}"].value
        if _n and isinstance(_v, (int, float)):
            cook_min[_n] = math.ceil(_v - 1e-9)

# 印刷用: 表示切替どおりの絞り込みと並び(作る順＝保持時間の長い順、未入力は最後)
if a.view == VIEW_HI:
    included = [i for i in range(NP) if pri[i] != PRI_LO]
elif a.view == VIEW_LO:
    included = [i for i in range(NP) if pri[i] == PRI_LO]
else:
    included = list(range(NP))
order = [i + 1 for i in sorted(included, key=lambda i: (ht[i] is None, -(ht[i] or 0), i))]
for k in range(N_SLOTS):
    r = 8 + k
    if k < len(order):
        idx = order[k]
        check(f"印刷用!B{r}(作る順)", pr[f"B{r}"].value, k + 1)
        check(f"印刷用!C{r}(商品名)", pr[f"C{r}"].value, PRODUCTS[idx - 1])
        check(f"印刷用!D{r}(開始目安=準備数計算Q)", as_serial(pr[f"D{r}"].value), as_serial(m[f"Q{ROW_M0 + idx - 1}"].value))
        check(f"印刷用!E{r}(調理の目安)", pr[f"E{r}"].value, cook_min.get(PRODUCTS[idx - 1], "—"))
        check(f"印刷用!F{r}(作るタイミング)", pr[f"F{r}"].value, pri[idx - 1])
        check(f"印刷用!G{r}(作る数=準備数計算G)", pr[f"G{r}"].value, m[f"G{ROW_M0 + idx - 1}"].value)
        check(f"印刷用!H{r}(☐)", pr[f"H{r}"].value, "☐")
    else:
        for c in "BCDEFGH":
            check(f"印刷用!{c}{r}(空)", pr[f"{c}{r}"].value in (None, ""), True)
if peak is not None:
    check("印刷用!B4にピーク表示", "ピーク " in (pr["B4"].value or ""), True)

# ---- 調理時間シート(あれば): 作る数の連動、回数=ROUNDUP(作る数/最大)、所要時間=(回数×最大個数の時間+(回数-1)×間隔)/60、
#      保持時間の連動、単純合計、機器ごとの合計
if "調理時間" in wb.sheetnames:
    import add_cooking_sheet as CK
    ck = wb["調理時間"]
    g_by_name = {m[f"C{ROW_M0 + i}"].value: m[f"G{ROW_M0 + i}"].value for i in range(N_SLOTS) if m[f"C{ROW_M0 + i}"].value}
    hold_by_name = {pd[f"B{14 + i}"].value: pd[f"E{14 + i}"].value for i in range(N_SLOTS) if pd[f"B{14 + i}"].value}
    tot_n, tot_p, seen, by_mach = 0, 0.0, 0, {}
    first, lastrow = CK.ROW_C0, CK.ROW_C0 + CK.N_ROWS - 1
    for r in range(first, lastrow + 1):
        k = r - first
        name = ck[f"{CK.COL_NAME}{r}"].value
        if k < NP:
            check(f"調理時間!{CK.COL_NAME}{r}(登録商品の順)", name, PRODUCTS[k])
        if not name:
            for c in (CK.COL_NEED, CK.COL_RUNS, CK.COL_MIN, CK.COL_HOLD):
                check(f"調理時間!{c}{r}(空)", ck[f"{c}{r}"].value in (None, ""), True)
            continue
        want_n = g_by_name.get(name, "未登録")
        if want_n in (None, ""):
            want_n = ""
        check(f"調理時間!{CK.COL_NEED}{r}(作る数)", ck[f"{CK.COL_NEED}{r}"].value in (None, "") if want_n == "" else ck[f"{CK.COL_NEED}{r}"].value,
              True if want_n == "" else want_n)
        hv = hold_by_name.get(name)
        check(f"調理時間!{CK.COL_HOLD}{r}(保持時間)", ck[f"{CK.COL_HOLD}{r}"].value, hv if isinstance(hv, (int, float)) else "—")
        mx = ck[f"{CK.COL_MAX}{r}"].value
        gap = ck[f"{CK.COL_GAP}{r}"].value
        gap = gap if isinstance(gap, (int, float)) else 0
        if isinstance(want_n, (int, float)):
            tot_n += want_n
            if not isinstance(mx, (int, float)):
                check(f"調理時間!{CK.COL_RUNS}{r}(最大未入力)", ck[f"{CK.COL_RUNS}{r}"].value, "⚠ 最大未入力")
                check(f"調理時間!{CK.COL_MIN}{r}(空)", ck[f"{CK.COL_MIN}{r}"].value in (None, ""), True)
                continue
            if mx <= 0 or mx > CK.COUNTS:
                check(f"調理時間!{CK.COL_RUNS}{r}(最大の⚠)", str(ck[f"{CK.COL_RUNS}{r}"].value).startswith("⚠"), True)
                check(f"調理時間!{CK.COL_MIN}{r}(空)", ck[f"{CK.COL_MIN}{r}"].value in (None, ""), True)
                continue
            want_o = math.ceil(want_n / mx - 1e-9)
            check(f"調理時間!{CK.COL_RUNS}{r}(回数)", ck[f"{CK.COL_RUNS}{r}"].value, want_o)
            t = ck.cell(row=r, column=CK.C_T0 - 1 + int(mx)).value
            if isinstance(t, (int, float)):
                want_p = (want_o * t + max(0, want_o - 1) * gap) / 60
                check(f"調理時間!{CK.COL_MIN}{r}(所要時間)", ck[f"{CK.COL_MIN}{r}"].value, want_p, tol=1e-6)
                tot_p += want_p
                seen += 1
                mach = ck[f"{CK.COL_MACH}{r}"].value
                if mach:
                    by_mach[mach] = by_mach.get(mach, 0) + want_p
            else:
                check(f"調理時間!{CK.COL_MIN}{r}(不可⚠)", ck[f"{CK.COL_MIN}{r}"].value, "⚠ その個数は不可／時間なし")
        else:
            check(f"調理時間!{CK.COL_RUNS}{r}(空)", ck[f"{CK.COL_RUNS}{r}"].value in (None, ""), True)
            check(f"調理時間!{CK.COL_MIN}{r}(空)", ck[f"{CK.COL_MIN}{r}"].value in (None, ""), True)
    tr = lastrow + 1
    check(f"調理時間!{CK.COL_NEED}{tr}(作る数の合計)", ck[f"{CK.COL_NEED}{tr}"].value, tot_n if tot_n else "", tol=0)
    if seen:
        check(f"調理時間!{CK.COL_MIN}{tr}(単純合計)", ck[f"{CK.COL_MIN}{tr}"].value, tot_p, tol=1e-6)
    else:
        check(f"調理時間!{CK.COL_MIN}{tr}(単純合計=空)", ck[f"{CK.COL_MIN}{tr}"].value in (None, ""), True)
    # 機器ごとの合計 = 同じ機器の所要時間の合計 ÷ 台数(1未満・未入力は1台扱い)
    r = tr + 1
    while r < tr + 8 and ck[f"{CK.COL_MACH}{r}"].value != "調理機器":
        r += 1
    check("調理時間!機器ごとの合計の見出し", ck[f"{CK.COL_MACH}{r}"].value, "調理機器")
    r += 1
    while ck[f"{CK.COL_MACH}{r}"].value:
        mach = ck[f"{CK.COL_MACH}{r}"].value
        units = ck[f"{CK.COL_MAX}{r}"].value
        units = units if isinstance(units, (int, float)) and units >= 1 else 1
        want = by_mach.get(mach)
        want = "—" if want is None else want / units
        check(f"調理時間!{CK.COL_MIN}{r}(機器合計 {mach[:12]}/{units}台)", ck[f"{CK.COL_MIN}{r}"].value, want,
              tol=1e-6 if want != "—" else 0)
        r += 1

warn = m["B9"].value or ""
# 「※」で始まる節は通知(計算は継続)なので、警告有無の判定から除く。「⚠」だけを警告として見る
warn_core = _re.sub(r"※[^⚠※]*", "", warn)
if a.expect_warn or a.rate == "blank" or int(a.rate) <= 0:
    for w in a.expect_warn + (["事前準備率が未入力か0以下"] if (a.rate == "blank" or int(a.rate) <= 0) else []):
        check(f"B9警告[{w}]", w in warn, True)
else:
    if warn_core.strip():
        errors.append(f"NG B9警告が出ている: {warn!r}")
    for ref in ("B7", "B11"):                  # 期間データの貼付状態にも⚠が無いこと(古いデータの通知は可)
        st = _re.sub(r"⚠ 古いデータの可能性（[^）]*）", "", pd[ref].value or "")
        check(f"期間データ!{ref}(⚠なし)", "⚠" in st, False)
check("印刷用!B6=B9", (pr["B6"].value or "") == warn, True)

if errors:
    print(f"FAILED ({len(errors)}):")
    for e in errors[:40]:
        print(" ", e)
    sys.exit(1)
print(f"ALL OK v2 (products={NP}, view={a.view}, rate={a.rate}, peak={a.peak_start or '-'}, "
      f"mso={'exact' if exact_mso else a.mso}, warn={a.expect_warn})")
