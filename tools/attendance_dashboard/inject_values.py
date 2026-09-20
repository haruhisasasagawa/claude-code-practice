#!/usr/bin/env python3
"""LibreOffice で再計算したブックの計算結果を、openpyxl が生成したブックにキャッシュ値として書き戻す.

openpyxl が書いた数式セルには計算結果（キャッシュ値）が無いため、Excel の保護ビューやプレビューでは
空欄に見える。同じ数式のブックを LibreOffice で再計算したファイルから値だけを取り出し、
生成ブックの XML に <v> として埋め込む（書式・グラフ・入力規則は生成ブックのまま）。

使い方:
    python inject_values.py 生成ブック.xlsx 再計算済み.xlsx 出力.xlsx
"""
import datetime as dt
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_RNS = "http://schemas.openxmlformats.org/package/2006/relationships"


def sheet_paths(zf):
    """シート名 → zip 内のパス."""
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid2target = {r.get("Id"): r.get("Target") for r in rels.iter(f"{{{PKG_RNS}}}Relationship")}
    out = {}
    for s in wb.iter(f"{{{NS}}}sheet"):
        target = rid2target[s.get(f"{{{RNS}}}id")].lstrip("/")
        out[s.get("name")] = target if target.startswith("xl/") else "xl/" + target
    return out


def cached_values(recalc_path):
    """再計算済みブックから {シート名: {座標: 値}} を読む（読み取り専用モードで省メモリ）."""
    wb = load_workbook(recalc_path, data_only=True, read_only=True)
    values = {}
    for ws in wb.worksheets:
        d = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    d[c.coordinate] = c.value
        values[ws.title] = d
    wb.close()
    return values


def inject_sheet(xml_bytes, vals):
    root = etree.fromstring(xml_bytes)
    n_set = 0
    for c in root.iter(f"{{{NS}}}c"):
        f = c.find(f"{{{NS}}}f")
        if f is None:
            continue
        v = vals.get(c.get("r"))
        old = c.find(f"{{{NS}}}v")
        if old is not None:
            c.remove(old)
        if v is None or v == "":
            c.set("t", "str")
            etree.SubElement(c, f"{{{NS}}}v").text = ""
            continue
        if isinstance(v, bool):
            c.set("t", "b"); text = "1" if v else "0"
        elif isinstance(v, (int, float)):
            c.attrib.pop("t", None); text = repr(float(v)) if isinstance(v, float) else str(v)
        elif isinstance(v, (dt.datetime, dt.date, dt.time, dt.timedelta)):
            c.attrib.pop("t", None); text = repr(float(to_excel(v)))
        elif isinstance(v, str) and v.startswith("#"):
            c.set("t", "e"); text = v
        else:
            c.set("t", "str"); text = str(v)
        ve = etree.SubElement(c, f"{{{NS}}}v")
        ve.text = text
        n_set += 1
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), n_set


def main(src, recalc, dst):
    values = cached_values(recalc)
    with zipfile.ZipFile(src) as zin:
        paths = sheet_paths(zin)
        path2sheet = {p: n for n, p in paths.items()}
        tmp = Path(tempfile.mkdtemp()) / "out.xlsx"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename in path2sheet:
                    name = path2sheet[item.filename]
                    data, n = inject_sheet(data, values.get(name, {}))
                    print(f"{name}: {n} values")
                elif item.filename == "xl/workbook.xml":
                    # 値を埋めたので、開いたときの全再計算は不要（現行の calcId にして fullCalcOnLoad を外す）
                    data = re.sub(rb'<calcPr[^>]*/>', b'<calcPr calcId="191029"/>', data)
                zout.writestr(item, data)
    shutil.move(str(tmp), dst)
    print("written", dst)


if __name__ == "__main__":
    main(*sys.argv[1:4])
