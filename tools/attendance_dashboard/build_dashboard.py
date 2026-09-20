#!/usr/bin/env python3
"""TOHOシネマズ新宿 勤務実績ダッシュボード（Excel）生成スクリプト.

シェアフルシフトから出力した月次シフトCSVを「CSV_1〜CSV_6」シートに貼り付けるだけで、
スタッフごとの出勤日数・出勤率・欠勤率などをダッシュボードに表示するExcelブックを生成する。
マクロは使わず、すべてワークシート関数で計算する（Excel 2016以降 / Microsoft 365 想定）。

使い方:
    python build_dashboard.py 出力.xlsx [--csv 8月.csv 9月.csv ...] [--maxrows 6000] [--select 氏名]
"""
import argparse
import csv
import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

# ---------------------------------------------------------------- constants
NSHEETS = 6
MAXSTAFF = 400            # 半年で扱えるスタッフ数の上限
STACK = NSHEETS * MAXSTAFF
FONT = "Yu Gothic"

S_HOWTO, S_DASH, S_LIST, S_SET, S_ROSTER = "使い方", "ダッシュボード", "スタッフ一覧", "設定", "名簿"
S_CSV = "CSV_{}"
S_CALC = "計算{}"

CSV_HEADERS = [
    "募集シフトの日付", "募集店舗", "応募ステータス", "募集シフトの開始時間", "募集シフトの終了時間",
    "相談応募の開始時間", "相談応募の終了時間", "アサインの開始時間", "アサインの終了時間",
    "変更後の開始時間", "変更後の終了時間", "募集シフトの職種", "募集シフトの資格", "応募者の名前",
    "応募者の従業員番号", "応募者の所属店舗", "応募者の電話番号", "応募者の職種", "応募者の資格", "時給",
    "休憩1開始時間", "休憩1終了時間", "休憩2開始時間", "休憩2終了時間", "休憩3開始時間", "休憩3終了時間",
    "連携ID", "勤務店舗コード", "所属店舗コード", "更新時間", "スケジュールID", "パターン名", "パターンコード", "勤務種別",
]
# 計算シートが列名で探す項目 → マッピングセル（計算シートの AB2〜AT2）
NEEDED = [
    ("AB", "募集シフトの日付"), ("AC", "応募ステータス"), ("AD", "募集シフトの開始時間"), ("AE", "募集シフトの終了時間"),
    ("AF", "変更後の開始時間"), ("AG", "変更後の終了時間"), ("AH", "募集シフトの職種"), ("AI", "応募者の名前"),
    ("AJ", "応募者の従業員番号"), ("AK", "応募者の職種"), ("AL", "応募者の資格"), ("AM", "休憩1開始時間"),
    ("AN", "休憩1終了時間"), ("AO", "休憩2開始時間"), ("AP", "休憩2終了時間"), ("AQ", "休憩3開始時間"),
    ("AR", "休憩3終了時間"), ("AS", "更新時間"), ("AT", "勤務種別"),
]
MAP = {name: colref for colref, name in NEEDED}

JOB_CATS = [("02コンセ", "コンセ"), ("03フロア", "フロア"), ("04ストア", "ストア"), ("05オフィス", "オフィス"),
            ("06トレーナー", "トレーナー"), ("07トレーニー", "トレーニー")]
DEPTS = [("02コンセ", "コンセ"), ("03フロア", "フロア"), ("04ストア", "ストア"), ("05オフィス", "オフィス")]
WEEKDAYS = ["日", "月", "火", "水", "木", "金", "土"]
REST_CATS = [("事前", "事前（2日以上前）"), ("前日", "前日"), ("当日", "当日（＝欠勤）"), ("事後", "シフト日より後")]

# 配色（dataviz パレット準拠）
C_BLUE, C_BLUE_D, C_RED = "2A78D6", "1C5CAB", "D03B3B"
C_GOOD, C_WARN, C_CRIT = "0CA30C", "FAB219", "D03B3B"
C_GOOD_TXT, C_WARN_TXT = "006300", "9A6400"
C_INK, C_INK2, C_MUTED, C_GRID, C_SURF, C_PAGE, C_DARK = "0B0B0B", "52514E", "898781", "E1E0D9", "FFFFFF", "F3F3F0", "1A1A19"
C_GRAY_BAR = "C3C2B7"
CAT_COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7"]
C_INPUT_FILL = "FFF9C4"

thin = Side(style="thin", color=C_GRID)
hair = Side(style="hair", color=C_GRID)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def q(name):
    return f"'{name}'"


def font(size=10, bold=False, color=C_INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def style(cell, *, size=10, bold=False, color=C_INK, bg=None, align=None, valign="center", wrap=False,
          fmt=None, border=None, italic=False):
    cell.font = font(size, bold, color, italic)
    if bg:
        cell.fill = fill(bg)
    cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border


def box_range(ws, rng, side=thin):
    """範囲の外枠だけに罫線を引く."""
    cells = list(ws[rng])
    top, bottom = cells[0], cells[-1]
    for c in top:
        c.border = Border(top=side, left=c.border.left, right=c.border.right, bottom=c.border.bottom)
    for c in bottom:
        c.border = Border(bottom=side, left=c.border.left, right=c.border.right, top=c.border.top)
    for row in cells:
        a, b = row[0], row[-1]
        a.border = Border(left=side, top=a.border.top, bottom=a.border.bottom, right=a.border.right)
        b.border = Border(right=side, top=b.border.top, bottom=b.border.bottom, left=b.border.left)


def fill_range(ws, rng, color):
    for row in ws[rng]:
        for c in row:
            c.fill = fill(color)


# ---------------------------------------------------------------- CSV → 貼付シート（Excel の貼り付け結果を再現）
_DATE = re.compile(r"^(\d{4})/(\d{1,2})/(\d{1,2})$")
_TIME = re.compile(r"^(\d{1,2}):(\d{2})$")
_NUM = re.compile(r"^-?\d+(\.\d+)?$")


def paste_value(s):
    """Excel に CSV を貼り付けたときの自動変換を再現する（日付・時刻・数値）."""
    if s is None or s == "":
        return None, None
    m = _DATE.match(s)
    if m:
        return dt.date(int(m[1]), int(m[2]), int(m[3])), "yyyy/mm/dd"
    m = _TIME.match(s)
    if m:
        return (int(m[1]) * 60 + int(m[2])) / 1440.0, "[h]:mm"
    if _NUM.match(s):
        return (int(s) if "." not in s else float(s)), None
    return s, None


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


# ---------------------------------------------------------------- シート: CSV_k
def build_csv_sheet(wb, k, rows=None, paste_mode="excel"):
    ws = wb.create_sheet(S_CSV.format(k))
    ws.sheet_properties.tabColor = "1BAF7A"
    for j, h in enumerate(CSV_HEADERS, 1):
        c = ws.cell(row=1, column=j, value=h)
        style(c, size=9, color=C_MUTED, bg="F3F3F0")
        ws.column_dimensions[L(j)].width = 13
    ws["A1"].comment = Comment(
        "シェアフルシフトのCSVをExcelで開き、ヘッダー行ごと全体をコピーして、このシートのA1セルに貼り付けてください。"
        "（先に貼ってあるデータは削除してから貼り付け）", "dashboard")
    if rows:
        for i, row in enumerate(rows[1:], 2):
            for j, s in enumerate(row, 1):
                v, fmt = paste_value(s) if paste_mode == "excel" else ((s or None), None)
                if v is None:
                    continue
                c = ws.cell(row=i, column=j, value=v)
                if fmt:
                    c.number_format = fmt
        for j, h in enumerate(rows[0], 1):
            ws.cell(row=1, column=j, value=h)
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- シート: 計算k
CALC_COLS = [
    ("A", "番号"), ("B", "日付"), ("C", "確定"), ("D", "種別"), ("E", "勤務行"), ("F", "更新日(JST)"),
    ("G", "欠勤行"), ("H", "出勤日(初回)"), ("I", "欠勤日(初回)"), ("J", "開始"), ("K", "終了"), ("L", "休憩"),
    ("M", "実働h"), ("N", "深夜h"), ("O", "曜日"), ("P", "休み区分"), ("Q", "却下"), ("R", "職種"),
    ("S", "初出"), ("T", "番号累積"), ("U", "休み日(初回)"), ("V", "更新時刻(JST)"), ("W", "選択者欠勤連番"),
    ("X", "元開始"), ("Y", "元終了"),
]


def build_calc_sheet(wb, k, maxrows):
    ws = wb.create_sheet(S_CALC.format(k))
    ws.sheet_properties.tabColor = "898781"
    csvs = q(S_CSV.format(k))
    last = maxrows + 1

    for colref, title in CALC_COLS:
        style(ws[f"{colref}1"], size=9, bold=True, color="FFFFFF", bg=C_INK2)
        ws.column_dimensions[colref].width = 9
    ws.column_dimensions["B"].width = 11
    ws.column_dimensions["F"].width = 11
    ws.column_dimensions["R"].width = 12

    # 列名マッピング（列順が変わっても動くように、ヘッダー名で列位置を探す）
    style(ws["AA1"], size=9, bold=True, color="FFFFFF", bg=C_INK2)
    ws["AA1"] = "列チェック"
    ws["AA2"] = f"=AND(COUNT($AB$2:$AT$2)={len(NEEDED)},MIN($AB$2:$AT$2)>0)"
    for colref, name in NEEDED:
        ws[f"{colref}1"] = name
        style(ws[f"{colref}1"], size=8, color=C_INK2, bg="F3F3F0")
        ws[f"{colref}2"] = f'=IFERROR(MATCH({colref}$1,{csvs}!$1:$1,0),IFERROR(MATCH("*"&{colref}$1&"*",{csvs}!$1:$1,0),0))'
        ws.column_dimensions[colref].width = 8
    ws["AA4"] = "列AB〜ATは貼付シートの列名→列番号の対応表、Z6〜AA12は貼付チェック用のセルです。数式で使うので編集しないでください。"
    style(ws["AA4"], size=8, color=C_MUTED)
    dash = q(S_DASH)
    brng = f"$B$2:$B${last}"
    arng = f"$A$2:$A${last}"
    checks = [
        (6, "対象月の初日（日付の中央値）", f'=IFERROR(DATE(YEAR(MEDIAN({brng})),MONTH(MEDIAN({brng})),1),"")'),
        (7, "対象月以外の日付の行数", f'=IF($AA$6="",0,SUMPRODUCT(({brng}<>"")*(({brng}<$AA$6)+({brng}>=EDATE($AA$6,1)))))'),
        (8, "日付を読み取れない行数", f'=SUMPRODUCT(({arng}<>"")*({brng}=""))'),
        (9, "番号が数値でない行数", f'=SUMPRODUCT(({arng}<>"")*ISTEXT({arng}))'),
        (11, "選択スタッフの欠勤日数（このシート）", f'=IF({dash}!$AC$3="",0,COUNTIFS({arng},{dash}!$AC$3,$I$2:$I${last},1))'),
        (12, "欠勤連番のオフセット", "=" + ("+".join(f"{q(S_CALC.format(j))}!$AA$11" for j in range(1, k)) if k > 1 else "0")),
    ]
    for r, label, formula in checks:
        ws[f"Z{r}"], ws[f"AA{r}"] = label, formula
        style(ws[f"Z{r}"], size=8, color=C_INK2)
    ws["AA6"].number_format = "yyyy/mm/dd"
    ws.column_dimensions["Z"].width = 30

    def idx(name):
        return f"INDEX({csvs}!$A:$AZ,ROW(),{MAP[name]}$2)"

    def tconv(name):
        x = idx(name)
        return f"IF(ISNUMBER({x}),{x},IFERROR(VALUE({x}),0))"

    num, date_, status, kind, upd = idx("応募者の従業員番号"), idx("募集シフトの日付"), idx("応募ステータス"), idx("勤務種別"), idx("更新時間")
    brk = "+".join(f"({tconv(f'休憩{i}終了時間')}-{tconv(f'休憩{i}開始時間')})" for i in (1, 2, 3))
    sets = q(S_SET)

    for r in range(2, last + 1):
        f = {
            "A": f'=IF(NOT($AA$2),"",IF({num}="","",IF(ISNUMBER({num}),{num},IFERROR(VALUE({num}),{num}))))',
            "B": f'=IF($A{r}="","",IF(ISNUMBER({date_}),{date_},IFERROR(DATEVALUE({date_}),IFERROR(DATEVALUE(SUBSTITUTE({date_},"/","-")),""))))',
            "C": f'=IF($A{r}="","",IF(ISNUMBER(SEARCH("確定",{status})),1,0))',
            "D": f'=IF($A{r}="","",IFERROR({kind}*1,0))',
            "E": f'=IF($A{r}="","",IF(AND($C{r}=1,$D{r}=1),1,0))',
            "F": f'=IF($A{r}="","",IFERROR(INT({upd}*1/86400000+25569+9/24),""))',
            "G": f'=IF($A{r}="","",IF(AND($C{r}=1,$D{r}=4,$B{r}<>"",$F{r}=$B{r}),1,0))',
            "H": f'=IF($A{r}="","",IF($E{r}=1,IF(COUNTIFS($A$2:$A{r},$A{r},$B$2:$B{r},$B{r},$E$2:$E{r},1)=1,1,0),0))',
            "I": (f'=IF($A{r}="","",IF($G{r}=1,IF(COUNTIFS($A$2:$A${last},$A{r},$B$2:$B${last},$B{r},$E$2:$E${last},1)>0,0,'
                  f'IF(COUNTIFS($A$2:$A{r},$A{r},$B$2:$B{r},$B{r},$G$2:$G{r},1)=1,1,0)),0))'),
            "J": f'=IF($E{r}<>1,"",IF({idx("変更後の開始時間")}="",{tconv("募集シフトの開始時間")},{tconv("変更後の開始時間")}))',
            "K": f'=IF($E{r}<>1,"",IF({idx("変更後の終了時間")}="",{tconv("募集シフトの終了時間")},{tconv("変更後の終了時間")}))',
            "L": f'=IF($E{r}<>1,"",{brk})',
            "M": f'=IF($E{r}<>1,"",($K{r}-$J{r}-$L{r})*24)',
            "N": f'=IF($E{r}<>1,"",MAX(0,MIN($K{r},{sets}!$B$8)-MAX($J{r},{sets}!$B$7))*24)',
            "O": f'=IF($A{r}="","",IF($B{r}="","",WEEKDAY($B{r})))',
            "P": (f'=IF($A{r}="","",IF(AND($C{r}=1,$D{r}=4),IF(OR($F{r}="",$B{r}=""),"不明",'
                  f'IF($F{r}>$B{r},"事後",IF($F{r}=$B{r},"当日",IF($F{r}=$B{r}-1,"前日","事前")))),""))'),
            "Q": f'=IF($A{r}="","",IF(ISNUMBER(SEARCH("却下",{status})),1,0))',
            "R": f'=IF($A{r}="","",{idx("募集シフトの職種")}&"")',
            "V": f'=IF($A{r}="","",IFERROR({upd}*1/86400000+25569+9/24,""))',
            "W": (f'=IF($A{r}="","",IF(AND($I{r}=1,$A{r}={dash}!$AC$3),'
                  f'COUNTIFS($A$2:$A${last},$A{r},$I$2:$I${last},1,$B$2:$B${last},"<"&$B{r})+1+$AA$12,""))'),
            # 休み行の元シフト時間。5:00〜29:00 は「全日休み」の既定値なので元の時間が残っていないとみなす
            "X": (f'=IF($A{r}="","",IF(AND($C{r}=1,$D{r}=4),IF(AND(ROUND({tconv("募集シフトの開始時間")}*1440,0)=300,'
                  f'ROUND({tconv("募集シフトの終了時間")}*1440,0)=1740),"",{tconv("募集シフトの開始時間")}),""))'),
            "Y": f'=IF($A{r}="","",IF($X{r}="","",{tconv("募集シフトの終了時間")}))',
            "S": f'=IF($A{r}="","",IF(COUNTIF($A$2:$A{r},$A{r})=1,1,0))',
            "T": f'=N($T{r-1})+IF($S{r}=1,1,0)',
            "U": (f'=IF($A{r}="","",IF(AND($C{r}=1,$D{r}=4),'
                  f'IF(COUNTIFS($A$2:$A{r},$A{r},$B$2:$B{r},$B{r},$C$2:$C{r},1,$D$2:$D{r},4)=1,1,0),0))'),
        }
        for colref, formula in f.items():
            ws[f"{colref}{r}"] = formula
        ws[f"B{r}"].number_format = "yyyy/mm/dd"
        ws[f"F{r}"].number_format = "yyyy/mm/dd"
        ws[f"J{r}"].number_format = "[h]:mm"
        ws[f"K{r}"].number_format = "[h]:mm"
        ws[f"L{r}"].number_format = "[h]:mm"
        ws[f"M{r}"].number_format = "0.00"
        ws[f"N{r}"].number_format = "0.00"
        ws[f"V{r}"].number_format = "yyyy/mm/dd hh:mm"
        ws[f"X{r}"].number_format = "[h]:mm"
        ws[f"Y{r}"].number_format = "[h]:mm"
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- シート: 設定
def build_settings(wb):
    ws = wb.create_sheet(S_SET)
    ws.sheet_properties.tabColor = "EDA100"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    ws["A1"] = "設定（黄色のセルは変更できます）"
    style(ws["A1"], size=14, bold=True)

    rows = [
        (3, "出勤率 ◎良好 の基準（この値以上）", 0.95, "0%", "総合判定に使用。既定 95%"),
        (4, "出勤率 △注意 の基準（この値以上）", 0.90, "0%", "この値未満は「✕要改善」。既定 90%"),
        (5, "判定を保留する確定シフト日数（この日数未満は「参考値」）", 20, "0", "シフト日数が少ないと欠勤1日で出勤率が大きく動くため、判定は出さず出勤率のみ表示。既定 20日"),
        (7, "深夜勤務の開始時刻", 22 / 24, "[h]:mm", "深夜勤務時間の集計範囲（開始）"),
        (8, "深夜勤務の終了時刻", 29 / 24, "[h]:mm", "同（終了）。翌5時は 29:00 と入力"),
    ]
    for r, label, val, fmt, note in rows:
        ws[f"A{r}"] = label
        ws[f"B{r}"] = val
        style(ws[f"B{r}"], bg=C_INPUT_FILL, fmt=fmt, border=BOX, align="right")
        ws[f"C{r}"] = note
        style(ws[f"C{r}"], size=9, color=C_INK2)

    ws["A11"] = "職種の区分（ダッシュボードの「職種別 勤務時間」に使用）"
    style(ws["A11"], bold=True)
    ws["A12"], ws["B12"] = "CSVの職種名", "表示名"
    for c in ("A12", "B12"):
        style(ws[c], size=9, bold=True, color=C_INK, bg="E7E6E1")
    for i, (code, label) in enumerate(JOB_CATS):
        r = 13 + i
        ws[f"A{r}"], ws[f"B{r}"] = code, label
        style(ws[f"A{r}"], bg=C_INPUT_FILL, border=BOX)
        style(ws[f"B{r}"], bg=C_INPUT_FILL, border=BOX)
    ws["A19"], ws["B19"] = "（上記以外）", "その他"
    style(ws["B19"], bg=C_INPUT_FILL, border=BOX)

    ws["A22"] = "所属（CSVの「応募者の職種」）の一覧（スタッフ一覧の所属別集計に使用）"
    style(ws["A22"], bold=True)
    ws["A23"], ws["B23"] = "CSVの所属名", "表示名"
    for c in ("A23", "B23"):
        style(ws[c], size=9, bold=True, color=C_INK, bg="E7E6E1")
    for i, (code, label) in enumerate(DEPTS):
        r = 24 + i
        ws[f"A{r}"], ws[f"B{r}"] = code, label
        style(ws[f"A{r}"], bg=C_INPUT_FILL, border=BOX)
        style(ws[f"B{r}"], bg=C_INPUT_FILL, border=BOX)
    return ws


# ---------------------------------------------------------------- シート: 名簿（集計エンジン）
ROSTER_COLS = {
    "H": "名簿行", "I": "従業員番号", "J": "名前", "K": "所属", "L": "資格", "M": "出勤日数", "N": "欠勤日数",
    "O": "確定シフト日数", "P": "出勤率", "Q": "欠勤率", "R": "勤務時間", "S": "深夜時間", "T": "平均/日",
    "U": "却下回数", "V": "名前順キー", "W": "出勤率順位", "X": "出勤率順キー(降順)", "Y": "ワーストキー(昇順)",
    "AR": "在籍月数", "AS": "出勤率整数キー", "AW": "同率人数", "AT": "名前(名前順)", "AU": "番号(名前順)", "AV": "名簿行(名前順)",
}
MONTH_WORK = ["Z", "AA", "AB", "AC", "AD", "AE"]      # 月別出勤日数
MONTH_ABS = ["AF", "AG", "AH", "AI", "AJ", "AK"]      # 月別欠勤日数
MONTH_HRS = ["AL", "AM", "AN", "AO", "AP", "AQ"]      # 月別勤務時間
MONTH_HAS = ["BG", "BH", "BI", "BJ", "BK", "BL"]      # 月別データ有無（1/0）


def calc_rng(k, colref, maxrows):
    return f"{q(S_CALC.format(k))}!${colref}$2:${colref}${maxrows + 1}"


def sum6(colref, maxrows, key):
    return "+".join(f"SUMIFS({calc_rng(k, colref, maxrows)},{calc_rng(k, 'A', maxrows)},{key})" for k in range(1, NSHEETS + 1))


def build_roster(wb, maxrows):
    ws = wb.create_sheet(S_ROSTER)
    ws.sheet_properties.tabColor = "898781"
    last = maxrows + 1
    head = {"A": "一致行", "B": "番号", "C": "名前", "D": "所属", "E": "資格", "F": "初出", "G": "累積"}
    head.update(ROSTER_COLS)
    for i, c in enumerate(MONTH_WORK):
        head[c] = f"出勤日数 月{i + 1}"
    for i, c in enumerate(MONTH_ABS):
        head[c] = f"欠勤日数 月{i + 1}"
    for i, c in enumerate(MONTH_HRS):
        head[c] = f"勤務時間 月{i + 1}"
    for i, c in enumerate(MONTH_HAS):
        head[c] = f"データ有無 月{i + 1}"
    for colref, title in head.items():
        ws[f"{colref}1"] = title
        style(ws[f"{colref}1"], size=9, bold=True, color="FFFFFF", bg=C_INK2)
        ws.column_dimensions[colref].width = 10
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["J"].width = 22
    ws.column_dimensions["AT"].width = 22

    # --- 6シート分のユニーク番号を積み上げ（新しい月＝番号の大きいシートを先に置き、最新の名前を採用）
    r = 2
    for k in range(NSHEETS, 0, -1):
        calc, csvs = q(S_CALC.format(k)), q(S_CSV.format(k))
        for i in range(1, MAXSTAFF + 1):
            ws[f"A{r}"] = f'=IFERROR(MATCH({i},{calc}!$T$2:$T${last},0),"")'
            ws[f"B{r}"] = f'=IF($A{r}="","",INDEX({calc}!$A$2:$A${last},$A{r}))'
            ws[f"C{r}"] = f'=IF($A{r}="","",INDEX({csvs}!$A:$AZ,$A{r}+1,{calc}!$AI$2)&"")'
            ws[f"D{r}"] = f'=IF($A{r}="","",INDEX({csvs}!$A:$AZ,$A{r}+1,{calc}!$AK$2)&"")'
            ws[f"E{r}"] = f'=IF($A{r}="","",INDEX({csvs}!$A:$AZ,$A{r}+1,{calc}!$AL$2)&"")'
            ws[f"F{r}"] = f'=IF($B{r}="","",IF(COUNTIF($B$2:$B{r},$B{r})=1,1,0))'
            ws[f"G{r}"] = f'=N($G{r - 1})+IF($F{r}=1,1,0)'
            r += 1
    stack_last = r - 1

    # --- 和集合（マスター）と半年集計
    sets = q(S_SET)
    for n in range(1, MAXSTAFF + 1):
        r = n + 1
        key = f"$I{r}"
        ws[f"H{r}"] = f'=IFERROR(MATCH({n},$G$2:$G${stack_last},0),"")'
        ws[f"I{r}"] = f'=IF($H{r}="","",INDEX($B$2:$B${stack_last},$H{r}))'
        ws[f"J{r}"] = f'=IF($H{r}="","",INDEX($C$2:$C${stack_last},$H{r}))'
        ws[f"K{r}"] = f'=IF($H{r}="","",INDEX($D$2:$D${stack_last},$H{r}))'
        ws[f"L{r}"] = f'=IF($H{r}="","",INDEX($E$2:$E${stack_last},$H{r}))'
        ws[f"M{r}"] = f'=IF({key}="","",{sum6("H", maxrows, key)})'
        ws[f"N{r}"] = f'=IF({key}="","",{sum6("I", maxrows, key)})'
        ws[f"O{r}"] = f'=IF({key}="","",$M{r}+$N{r})'
        ws[f"P{r}"] = f'=IF({key}="","",IF($O{r}=0,"",$M{r}/$O{r}))'
        ws[f"Q{r}"] = f'=IF({key}="","",IF($O{r}=0,"",$N{r}/$O{r}))'
        ws[f"R{r}"] = f'=IF({key}="","",{sum6("M", maxrows, key)})'
        ws[f"S{r}"] = f'=IF({key}="","",{sum6("N", maxrows, key)})'
        ws[f"T{r}"] = f'=IF({key}="","",IF($M{r}=0,"",$R{r}/$M{r}))'
        ws[f"U{r}"] = f'=IF({key}="","",{sum6("Q", maxrows, key)})'
        ws[f"V{r}"] = f'=IF($J{r}="","",COUNTIFS($J$2:$J${MAXSTAFF + 1},"<"&$J{r},$J$2:$J${MAXSTAFF + 1},"?*")+COUNTIF($J$2:$J{r},$J{r}))'
        ws[f"AS{r}"] = f'=IF($P{r}="","",ROUND($P{r}*100000,0))'
        ws[f"W{r}"] = f'=IF($AS{r}="","",COUNTIF($AS$2:$AS${MAXSTAFF + 1},">"&$AS{r})+1)'
        ws[f"X{r}"] = f'=IF($AS{r}="","",$W{r}+COUNTIF($AS$2:$AS{r},$AS{r})-1)'
        ws[f"Y{r}"] = (f'=IF($AS{r}="","",IF($O{r}<{sets}!$B$5,"",'
                       f'COUNTIFS($AS$2:$AS${MAXSTAFF + 1},"<"&$AS{r},$O$2:$O${MAXSTAFF + 1},">="&{sets}!$B$5)'
                       f'+COUNTIFS($AS$2:$AS${MAXSTAFF + 1},$AS{r},$O$2:$O${MAXSTAFF + 1},">"&$O{r},$O$2:$O${MAXSTAFF + 1},">="&{sets}!$B$5)'
                       f'+COUNTIFS($AS$2:$AS{r},$AS{r},$O$2:$O{r},$O{r},$O$2:$O{r},">="&{sets}!$B$5)))')
        ws[f"AW{r}"] = f'=IF($AS{r}="","",COUNTIF($W$2:$W${MAXSTAFF + 1},$W{r}))'
        for k in range(1, NSHEETS + 1):
            ws[f"{MONTH_WORK[k - 1]}{r}"] = f'=IF({key}="","",SUMIFS({calc_rng(k, "H", maxrows)},{calc_rng(k, "A", maxrows)},{key}))'
            ws[f"{MONTH_ABS[k - 1]}{r}"] = f'=IF({key}="","",SUMIFS({calc_rng(k, "I", maxrows)},{calc_rng(k, "A", maxrows)},{key}))'
            ws[f"{MONTH_HRS[k - 1]}{r}"] = f'=IF({key}="","",SUMIFS({calc_rng(k, "M", maxrows)},{calc_rng(k, "A", maxrows)},{key}))'
        for k in range(1, NSHEETS + 1):
            ws[f"{MONTH_HAS[k - 1]}{r}"] = f'=IF({key}="","",IF(COUNTIF({calc_rng(k, "A", maxrows)},{key})>0,1,0))'
        ws[f"AR{r}"] = f'=IF({key}="","",SUM({MONTH_HAS[0]}{r}:{MONTH_HAS[-1]}{r}))'
        ws[f"AT{r}"] = f'=IFERROR(INDEX($J$2:$J${MAXSTAFF + 1},MATCH({n},$V$2:$V${MAXSTAFF + 1},0)),"")'
        ws[f"AU{r}"] = f'=IFERROR(INDEX($I$2:$I${MAXSTAFF + 1},MATCH({n},$V$2:$V${MAXSTAFF + 1},0)),"")'
        ws[f"AV{r}"] = f'=IFERROR(MATCH({n},$V$2:$V${MAXSTAFF + 1},0),"")'
        for c in ("P", "Q"):
            ws[f"{c}{r}"].number_format = "0.0%"
        for c in ("R", "S", "T"):
            ws[f"{c}{r}"].number_format = "0.0"

    # --- 全体・所属別サマリー
    ML = MAXSTAFF + 1
    ws["AX1"] = "全体サマリー"
    style(ws["AX1"], size=9, bold=True, color="FFFFFF", bg=C_INK2)
    ws.column_dimensions["AX"].width = 20
    ws.column_dimensions["AY"].width = 12
    summary = [
        (2, "スタッフ数", f"=SUMPRODUCT(--($I$2:$I${ML}<>\"\"))", "0"),
        (3, "総出勤日数", f"=SUM($M$2:$M${ML})", "0"),
        (4, "総欠勤日数", f"=SUM($N$2:$N${ML})", "0"),
        (5, "総確定シフト日数", f"=SUM($O$2:$O${ML})", "0"),
        (6, "全体出勤率", f'=IF($AY$5=0,"",$AY$3/$AY$5)', "0.0%"),
        (7, "全体欠勤率", f'=IF($AY$5=0,"",$AY$4/$AY$5)', "0.0%"),
        (8, "評価対象人数(確定日数>0)", f"=COUNT($P$2:$P${ML})", "0"),
        (9, "総勤務時間", f"=SUM($R$2:$R${ML})", "0.0"),
    ]
    for r, label, formula, fmt in summary:
        ws[f"AX{r}"], ws[f"AY{r}"] = label, formula
        ws[f"AY{r}"].number_format = fmt

    ws["AX12"] = "所属別"
    style(ws["AX12"], size=9, bold=True, color="FFFFFF", bg=C_INK2)
    for j, h in enumerate(["所属コード", "表示名", "人数", "出勤日数", "欠勤日数", "確定日数", "出勤率", "欠勤率"]):
        c = ws.cell(row=13, column=50 + j, value=h)   # AX=50
        style(c, size=9, bold=True, color="FFFFFF", bg=C_INK2)
    for i in range(len(DEPTS)):
        r = 14 + i
        ws[f"AX{r}"] = f"={sets}!$A${24 + i}"
        ws[f"AY{r}"] = f"={sets}!$B${24 + i}"
        ws[f"AZ{r}"] = f"=COUNTIF($K$2:$K${ML},$AX{r})"
        ws[f"BA{r}"] = f"=SUMIFS($M$2:$M${ML},$K$2:$K${ML},$AX{r})"
        ws[f"BB{r}"] = f"=SUMIFS($N$2:$N${ML},$K$2:$K${ML},$AX{r})"
        ws[f"BC{r}"] = f"=SUMIFS($O$2:$O${ML},$K$2:$K${ML},$AX{r})"
        ws[f"BD{r}"] = f'=IF($BC{r}=0,"",$BA{r}/$BC{r})'
        ws[f"BE{r}"] = f'=IF($BC{r}=0,"",$BB{r}/$BC{r})'
        ws[f"BD{r}"].number_format = "0.0%"
        ws[f"BE{r}"].number_format = "0.0%"

    # --- 月ラベル・貼付状況
    ws["AX20"] = "月ラベル / 貼付状況"
    style(ws["AX20"], size=9, bold=True, color="FFFFFF", bg=C_INK2)
    for j, h in enumerate(["シート", "月", "貼付行数", "状態"]):
        c = ws.cell(row=21, column=50 + j, value=h)
        style(c, size=9, bold=True, color="FFFFFF", bg=C_INK2)
    for k in range(1, NSHEETS + 1):
        r = 21 + k
        calc, csvs = q(S_CALC.format(k)), q(S_CSV.format(k))
        m1 = f"{calc}!$AA$6"
        ws[f"AX{r}"] = S_CSV.format(k)
        ws[f"AY{r}"] = f'=IF({m1}="","月{k}（未貼付）",YEAR({m1})&"年"&MONTH({m1})&"月")'
        ws[f"AZ{r}"] = f"=COUNTA({csvs}!$A$2:$A$200000)"
        ws[f"BA{r}"] = (f'=IF($AZ{r}=0,"未貼付",IF(NOT({calc}!$AA$2),"※ 列名が見つかりません（1行目にヘッダーを含めて貼り付けてください）",'
                        f'IF($AZ{r}>{maxrows},"※ {maxrows:,}行を超えています（超過分は集計されません）",'
                        f'IF({calc}!$AA$7>0,"※ 対象月以外の日付の行が "&{calc}!$AA$7&" 行あります（前のデータが残っている可能性。全選択→削除してから貼り直し）",'
                        f'IF({calc}!$AA$8+{calc}!$AA$9>0,"※ 読み取れない値があります（日付 "&{calc}!$AA$8&" 行／番号 "&{calc}!$AA$9&" 行）","OK")))))')
    ws["AX29"] = "対象期間"
    ws["AY29"] = (f'=IF(SUM($AZ$22:$AZ$27)=0,"（CSV未貼付）",'
                  f'INDEX($AY$22:$AY$27,MATCH(TRUE,INDEX($AZ$22:$AZ$27>0,0),0))&" 〜 "&'
                  f'INDEX($AY$22:$AY$27,SUMPRODUCT(MAX(($AZ$22:$AZ$27>0)*(ROW($AZ$22:$AZ$27)-21)))))')
    ws["AX30"] = "月順チェック"
    ws["AY30"] = "=IF(" + "+".join(
        f'IF(AND({q(S_CALC.format(k))}!$AA$6<>"",{q(S_CALC.format(k - 1))}!$AA$6<>"",{q(S_CALC.format(k))}!$AA$6<={q(S_CALC.format(k - 1))}!$AA$6),1,0)'
        for k in range(2, NSHEETS + 1)) + '>0,"※ CSV_1→CSV_6 が古い月から順に並んでいません（月別推移の並びと最新の名前の採用が崩れます）","")'
    ws["AX31"] = "警告あり"
    ws["AY31"] = '=IF(OR(COUNTIF($BA$22:$BA$27,"※*")>0,$AY$30<>""),1,0)'
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- シート: ダッシュボード
GRID_FIRST, GRID_LAST = 2, 24            # B〜X の23列をグリッドとして使う（A・Y は余白）
HC = "AC"                                # 内部計算セルの列（値）
HL = "AB"                                # 内部計算セルの列（ラベル）
ABS_ROWS = 12                            # 欠勤日付一覧の行数
WD_ORDER = [(2, "月"), (3, "火"), (4, "水"), (5, "木"), (6, "金"), (7, "土"), (1, "日")]   # WEEKDAY値と表示
C_HEAD = "E7E6E1"                        # 表見出しの薄いグレー
C_NAVY = "1F3A5F"


def build_dashboard(wb, maxrows, select=None):
    ws = wb.create_sheet(S_DASH)
    ws.sheet_properties.tabColor = C_NAVY
    ws.sheet_view.showGridLines = False
    ros, sets = q(S_ROSTER), q(S_SET)
    ML = MAXSTAFF + 1
    last = maxrows + 1
    ws.column_dimensions["A"].width = 2.5
    for c in range(GRID_FIRST, GRID_LAST + 1):
        ws.column_dimensions[L(c)].width = 4.6
    ws.column_dimensions["Y"].width = 2.5
    ws.column_dimensions["Z"].width = 2.5
    ws.column_dimensions["AA"].width = 3
    ws.column_dimensions[HL].width = 26
    ws.column_dimensions[HC].width = 14
    ws.column_dimensions["AD"].width = 12
    LAST_ROW = 82
    fill_range(ws, f"A1:Y{LAST_ROW}", "FFFFFF")

    def V(colref):
        return f'IF(${HC}$4="","",INDEX({ros}!${colref}$2:${colref}${ML},${HC}$4))'

    # ---- 内部計算セル（画面右外）
    ws[f"{HL}1"] = "内部計算（このブロックは触らないでください）"
    style(ws[f"{HL}1"], size=9, bold=True, color=C_MUTED)
    helpers = [
        (3, "選択スタッフの従業員番号", f'=IFERROR(INDEX({ros}!$AU$2:$AU${ML},MATCH($B$7,{ros}!$AT$2:$AT${ML},0)),"")'),
        (4, "名簿の行", f'=IF(${HC}$3="","",IFERROR(MATCH(${HC}$3,{ros}!$I$2:$I${ML},0),""))'),
        (5, "所属コード", f"={V('K')}"),
        (6, "出勤率", f"={V('P')}"),
        (7, "◎の基準", f"={sets}!$B$3"),
        (8, "△の基準", f"={sets}!$B$4"),
        (9, "判定保留の日数", f"={sets}!$B$5"),
        (10, "確定シフト日数", f"={V('O')}"),
    ]
    for r, label, formula in helpers:
        ws[f"{HL}{r}"], ws[f"{HC}{r}"] = label, formula
        style(ws[f"{HL}{r}"], size=9, color=C_INK2)
        style(ws[f"{HC}{r}"], size=9)
    ws[f"{HC}6"].number_format = "0.0%"
    sel = f"${HC}$3"

    def sum6sel(colref, crit_col=None, crit=None):
        parts = []
        for k in range(1, NSHEETS + 1):
            a = f"SUMIFS({calc_rng(k, colref, maxrows)},{calc_rng(k, 'A', maxrows)},{sel}"
            if crit_col:
                a += f",{calc_rng(k, crit_col, maxrows)},{crit}"
            parts.append(a + ")")
        return "+".join(parts)

    ws[f"{HL}12"], ws[f"{HC}12"] = "出勤", f"=N({V('M')})"
    ws[f"{HL}13"], ws[f"{HC}13"] = "欠勤", f"=N({V('N')})"
    for i in range(len(JOB_CATS)):
        r = 16 + i
        ws[f"{HL}{r}"] = f"={sets}!$B${13 + i}"
        ws[f"{HC}{r}"] = f'=IF({sel}="",0,{sum6sel("M", "R", f"{sets}!$A${13 + i}")})'
    ws[f"{HL}22"] = f"={sets}!$B$19"
    ws[f"{HC}22"] = f'=IF({sel}="",0,MAX(0,N({V("R")})-SUM(${HC}$16:${HC}$21)))'
    for i, (code, wd) in enumerate(WD_ORDER):
        r = 25 + i
        ws[f"{HL}{r}"] = wd
        ws[f"{HC}{r}"] = f'=IF({sel}="",0,{sum6sel("H", "O", code)})'
    ws[f"{HL}34"], ws[f"{HC}34"] = "本人", f"=N(${HC}$6)"
    ws[f"{HL}35"], ws[f"{HC}35"] = "全体平均", f"=N({ros}!$AY$6)"
    ws[f"{HL}36"] = f'=IF(${HC}$5="","同所属","同所属（"&IFERROR(INDEX({ros}!$AY$14:$AY$17,MATCH(${HC}$5,{ros}!$AX$14:$AX$17,0)),${HC}$5)&"）")'
    ws[f"{HC}36"] = f'=IFERROR(N(INDEX({ros}!$BD$14:$BD$17,MATCH(${HC}$5,{ros}!$AX$14:$AX$17,0))),0)'
    for r in (34, 35, 36):
        ws[f"{HC}{r}"].number_format = "0.0%"
    for i, (code, label) in enumerate(REST_CATS):
        r = 39 + i
        ws[f"{HL}{r}"] = label
        crit = '"' + code + '"'
        ws[f"{HC}{r}"] = f'=IF({sel}="",0,{sum6sel("U", "P", crit)})'
    ws[f"{HL}45"], ws[f"{HC}45"], ws["AD45"] = "月（グラフ用）", "出勤日数", "欠勤日数"
    for k in range(1, NSHEETS + 1):
        r = 45 + k
        ws[f"{HL}{r}"] = f'=IF({ros}!$AZ${21 + k}>0,MONTH({q(S_CALC.format(k))}!$AA$6)&"月","")'
        ws[f"{HC}{r}"] = f"=N({V(MONTH_WORK[k - 1])})"
        ws[f"AD{r}"] = f"=N({V(MONTH_ABS[k - 1])})"
        ws[f"{HC}{r}"].number_format = "0;;;"
        ws[f"AD{r}"].number_format = "0;;;"
    ws[f"{HL}54"], ws[f"{HC}54"] = "欠勤日数（一覧の件数）", "=" + "+".join(f"{q(S_CALC.format(k))}!$AA$11" for k in range(1, NSHEETS + 1))
    ws[f"{HL}55"], ws[f"{HC}55"] = "却下回数（参考）", f"=N({V('U')})"

    def section(row, text, a="B", b="X"):
        """見出し行: 太字＋下罫線."""
        ws[f"{a}{row}"] = text
        style(ws[f"{a}{row}"], size=10.5, bold=True, color=C_INK)
        for c in range(column_index_from_string(a), column_index_from_string(b) + 1):
            ws[f"{L(c)}{row}"].border = Border(bottom=Side(style="thin", color="9E9C95"))

    def head_cell(ref, text, align="center"):
        ws[ref] = text
        style(ws[ref], size=9, bold=True, color=C_INK, bg=C_HEAD, align=align, border=Border(bottom=thin, top=thin))

    # ---- タイトル
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 14
    ws.merge_cells("B2:N2")
    ws["B2"] = "勤務実績ダッシュボード"
    style(ws["B2"], size=16, bold=True, color=C_NAVY, align="left")
    ws.merge_cells("B3:N3")
    ws["B3"] = "TOHOシネマズ新宿　アルバイトスタッフ（シェアフルシフトのシフト実績より）"
    style(ws["B3"], size=9, color=C_INK2, align="left")
    ws.merge_cells("O2:X2")
    ws["O2"] = f'="対象期間　"&{ros}!$AY$29'
    style(ws["O2"], size=10.5, bold=True, color=C_INK, align="right")
    ws.merge_cells("O3:X3")
    ws["O3"] = (f'=IF({ros}!$AY$31=1,"※ 貼付データに問題があります。「使い方」の貼付状況をご確認ください",'
                f'IF({ros}!$AY$2="","","集計対象 "&{ros}!$AY$2&"名　　貼付済み "&COUNTIF({ros}!$BA$22:$BA$27,"OK")&"ヶ月"))')
    style(ws["O3"], size=9, color=C_INK2, align="right")
    ws.conditional_formatting.add("O3", FormulaRule(formula=['LEFT($O$3,1)="※"'], font=Font(name=FONT, bold=True, size=9, color=C_CRIT)))
    for c in range(GRID_FIRST, GRID_LAST + 1):
        ws[f"{L(c)}4"].border = Border(bottom=Side(style="medium", color=C_NAVY))

    # ---- スタッフ選択
    ws.row_dimensions[5].height = 14
    ws.row_dimensions[6].height = 15
    ws.row_dimensions[7].height = 22
    ws.row_dimensions[8].height = 22
    ws["B6"] = "スタッフ（▼から選択）"
    style(ws["B6"], size=9, color=C_INK2)
    ws.merge_cells("B7:H8")
    if select:
        ws["B7"] = select
    else:
        ws["B7"] = f'=IF({ros}!$AT$2="","← CSV_1 にデータを貼り付けてください",{ros}!$AT$2)'
    style(ws["B7"], size=14, bold=True, bg="FFFFFF", align="left")
    box_range(ws, "B7:H8", Side(style="medium", color=C_NAVY))
    dv = DataValidation(type="list", formula1="StaffNames", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("B7")

    def info(col_a, col_b, label, formula, fmt=None, size=11):
        ws[f"{col_a}6"] = label
        style(ws[f"{col_a}6"], size=9, color=C_INK2)
        ws.merge_cells(f"{col_a}7:{col_b}8")
        ws[f"{col_a}7"] = formula
        style(ws[f"{col_a}7"], size=size, bold=True, bg="FFFFFF", align="left", fmt=fmt)
        box_range(ws, f"{col_a}7:{col_b}8")

    info("J", "K", "従業員番号", f"={V('I')}", fmt="0")
    info("M", "N", "所属", f'=IF(${HC}$5="","",IFERROR(INDEX({ros}!$AY$14:$AY$17,MATCH(${HC}$5,{ros}!$AX$14:$AX$17,0)),${HC}$5))')
    info("P", "Q", "資格", f'=IF(${HC}$4="","",SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({V("L")},"01アルバイト","アルバイト"),"02サブリーダー","サブリーダー"),"03リーダー","リーダー"))', size=10)
    ws["S6"] = f'="総合判定（出勤率 "&TEXT(${HC}$7,"0%")&"以上◎／"&TEXT(${HC}$8,"0%")&"以上△）"'
    style(ws["S6"], size=9, color=C_INK2)
    ws.merge_cells("S7:X8")
    ws["S7"] = (f'=IF(${HC}$6="","－",IF(${HC}$10<${HC}$9,"参考値（確定シフト "&${HC}$10&"日）",'
                f'IF(${HC}$6>=${HC}$7,"◎ 良好",IF(${HC}$6>=${HC}$8,"△ 注意","✕ 要改善"))))')
    style(ws["S7"], size=13, bold=True, color=C_INK, bg="F3F3F0", align="center")
    box_range(ws, "S7:X8")
    badge_rules = [
        (f'AND(ISNUMBER(${HC}$6),${HC}$10<${HC}$9)', "EEEEEA", C_INK2),
        (f'AND(ISNUMBER(${HC}$6),${HC}$6>=${HC}$7)', "E3F4E3", C_GOOD_TXT),
        (f'AND(ISNUMBER(${HC}$6),${HC}$6>=${HC}$8)', "FFF1CC", C_WARN_TXT),
        (f'ISNUMBER(${HC}$6)', "F9DEDE", C_CRIT),
    ]
    for formula, bg, fg in badge_rules:
        ws.conditional_formatting.add("S7:X8", FormulaRule(formula=[formula], fill=fill(bg), font=Font(name=FONT, bold=True, color=fg, size=13), stopIfTrue=True))

    # ---- KPI（3列＋空き1列 ×6）
    ws.row_dimensions[9].height = 16
    for r, h in ((10, 16), (11, 22), (12, 22), (13, 16)):
        ws.row_dimensions[r].height = h
    all_rate = f"{ros}!$AY$6"
    tiles = [
        ("出勤日数", f"={V('M')}", '0"日"', f'=IF(${HC}$4="","","確定シフト "&{V("O")}&"日")'),
        ("出勤率", f"={V('P')}", "0.0%", f'=IF(${HC}$6="","",IF({all_rate}="","","全体平均 "&TEXT({all_rate},"0.0%")))'),
        ("欠勤日数", f"={V('N')}", '0"日"', '="当日に休みへ変更した日"'),
        ("欠勤率", f"={V('Q')}", "0.0%", f'=IF(${HC}$6="","",IF({ros}!$AY$7="","","全体平均 "&TEXT({ros}!$AY$7,"0.0%")))'),
        ("勤務時間", f"={V('R')}", '0.0"h"', f'=IF(${HC}$4="","",IF({V("T")}="","","1日あたり "&TEXT({V("T")},"0.0")&"h"))'),
        ("深夜勤務", f"={V('S')}", '0.0"h"', f'=IF(${HC}$4="","",IF(N({V("R")})=0,"","勤務時間の "&TEXT({V("S")}/{V("R")},"0%")))'),
    ]
    for i, (label, formula, fmt, sub) in enumerate(tiles):
        c0 = GRID_FIRST + i * 4
        a, b = L(c0), L(c0 + 2)
        ws.merge_cells(f"{a}10:{b}10")
        ws[f"{a}10"] = label
        style(ws[f"{a}10"], size=9, color=C_INK2, align="left")
        ws.merge_cells(f"{a}11:{b}12")
        ws[f"{a}11"] = formula
        style(ws[f"{a}11"], size=20, bold=True, align="left", fmt=fmt)
        ws.merge_cells(f"{a}13:{b}13")
        ws[f"{a}13"] = sub
        style(ws[f"{a}13"], size=8, color=C_MUTED, align="left")
        for c in range(c0, c0 + 3):
            ws[f"{L(c)}10"].border = Border(top=thin)
            ws[f"{L(c)}13"].border = Border(bottom=thin)
    rate_rules = [
        (f'AND(ISNUMBER(${HC}$6),${HC}$6>=${HC}$7)', C_GOOD_TXT),
        (f'AND(ISNUMBER(${HC}$6),${HC}$6>=${HC}$8)', C_WARN_TXT),
        (f'ISNUMBER(${HC}$6)', C_CRIT),
    ]
    for formula, color in rate_rules:
        ws.conditional_formatting.add("F11:H12", FormulaRule(formula=[formula], font=Font(name=FONT, bold=True, size=20, color=color), stopIfTrue=True))

    # ---- 一文サマリー
    ws.row_dimensions[14].height = 10
    ws.row_dimensions[15].height = 20
    ws.merge_cells("B15:X15")
    ws["B15"] = (f'=IF(${HC}$4="","",IF(N(${HC}$10)=0,"対象期間に確定シフトがありません。",'
                 f'"確定シフト "&${HC}$10&"日のうち出勤 "&{V("M")}&"日、当日欠勤 "&{V("N")}&"日。'
                 f'出勤率 "&TEXT(${HC}$6,"0.0%")&"（全体平均 "&TEXT(${HC}$35,"0.0%")&"、"&{ros}!$AY$8&"人中 "&{V("W")}&"位）"))')
    style(ws["B15"], size=10, color=C_INK, align="left")
    ws.row_dimensions[16].height = 14

    # ---- グラフ 1段目（7列＋空き1列 ×3）
    ws.row_dimensions[17].height = 18
    for r in range(18, 32):
        ws.row_dimensions[r].height = 15
    section(17, "出勤率", "B", "H")
    section(17, "月別の出勤日数・欠勤日数", "J", "P")
    section(17, "職種別の勤務時間", "R", "X")

    def gp(color):
        g = GraphicalProperties(solidFill=color)
        g.line.solidFill = "FFFFFF"
        return g

    def plain(ch):
        ch.graphical_properties = GraphicalProperties()
        ch.graphical_properties.line.noFill = True
        ch.width, ch.height = 6.8, 7.0
        return ch

    def doughnut(label_rng, data_rng, colors, hole=58):
        ch = DoughnutChart()
        ch.holeSize = hole
        ch.add_data(Reference(ws, range_string=f"{q(S_DASH)}!{data_rng}"), titles_from_data=False)
        ch.set_categories(Reference(ws, range_string=f"{q(S_DASH)}!{label_rng}"))
        s = ch.series[0]
        for i, color in enumerate(colors):
            pt = DataPoint(idx=i)
            pt.graphicalProperties = gp(color)
            s.dPt.append(pt)
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ch.dataLabels.showVal = False
        ch.dataLabels.showCatName = False
        ch.dataLabels.showSerName = False
        ch.dataLabels.showLeaderLines = False
        ch.dataLabels.showLegendKey = False
        ch.dataLabels.numFmt = "0%;;;"
        ch.legend.position = "b"
        return plain(ch)

    ws.add_chart(doughnut(f"${HL}$12:${HL}$13", f"${HC}$12:${HC}$13", [C_BLUE, C_RED], hole=60), "B18")
    ws.add_chart(doughnut(f"${HL}$16:${HL}$22", f"${HC}$16:${HC}$22", CAT_COLORS), "R18")

    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 60
    ch.add_data(Reference(ws, range_string=f"{q(S_DASH)}!${HC}$45:${HC}$51"), titles_from_data=True)
    ch.add_data(Reference(ws, range_string=f"{q(S_DASH)}!$AD$45:$AD$51"), titles_from_data=True)
    ch.set_categories(Reference(ws, range_string=f"{q(S_DASH)}!${HL}$46:${HL}$51"))
    ch.series[0].graphicalProperties = gp(C_BLUE)
    ch.series[1].graphicalProperties = gp(C_RED)
    for s_ in ch.series:
        s_.dLbls = DataLabelList()
        s_.dLbls.showVal = True
        s_.dLbls.showSerName = False
        s_.dLbls.showCatName = False
        s_.dLbls.showLegendKey = False
        s_.dLbls.numFmt = "0;;;"
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_GRID))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.scaling.min = 0
    ch.y_axis.number_format = "0"
    ch.legend.position = "b"
    ws.add_chart(plain(ch), "J18")

    # ---- グラフ 2段目
    ws.row_dimensions[32].height = 14
    ws.row_dimensions[33].height = 18
    for r in range(34, 48):
        ws.row_dimensions[r].height = 15
    section(33, "曜日別の出勤日数", "B", "H")
    section(33, "出勤率の比較", "J", "P")
    section(33, "休みへの変更のタイミング", "R", "X")

    ch = BarChart()
    ch.type, ch.gapWidth = "col", 50
    ch.add_data(Reference(ws, range_string=f"{q(S_DASH)}!${HC}$25:${HC}$31"), titles_from_data=False)
    ch.set_categories(Reference(ws, range_string=f"{q(S_DASH)}!${HL}$25:${HL}$31"))
    ch.series[0].graphicalProperties = gp(C_BLUE)
    for i in (5, 6):
        pt = DataPoint(idx=i)
        pt.graphicalProperties = gp(C_BLUE_D)
        ch.series[0].dPt.append(pt)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = False
    ch.dataLabels.showCatName = False
    ch.dataLabels.showLegendKey = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_GRID))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.scaling.min = 0
    ch.y_axis.number_format = "0"
    ch.legend = None
    ws.add_chart(plain(ch), "B34")

    ch = BarChart()
    ch.type, ch.gapWidth = "bar", 50
    ch.add_data(Reference(ws, range_string=f"{q(S_DASH)}!${HC}$34:${HC}$36"), titles_from_data=False)
    ch.set_categories(Reference(ws, range_string=f"{q(S_DASH)}!${HL}$34:${HL}$36"))
    ch.series[0].graphicalProperties = gp(C_GRAY_BAR)
    pt = DataPoint(idx=0)
    pt.graphicalProperties = gp(C_BLUE)
    ch.series[0].dPt.append(pt)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = False
    ch.dataLabels.showCatName = False
    ch.dataLabels.showLegendKey = False
    ch.dataLabels.numFmt = "0.0%"
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
    ch.y_axis.number_format = "0%"
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=C_GRID))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend = None
    ws.add_chart(plain(ch), "J34")

    ws.row_dimensions[35].height = 18
    ws.merge_cells("R35:V35")
    head_cell("R35", "区分", align="left")
    ws.merge_cells("W35:X35")
    head_cell("W35", "日数")
    for i, (code, label) in enumerate(REST_CATS):
        r = 36 + i * 2
        ws.merge_cells(f"R{r}:V{r + 1}")
        ws[f"R{r}"] = label
        style(ws[f"R{r}"], size=10, align="left")
        ws.merge_cells(f"W{r}:X{r + 1}")
        ws[f"W{r}"] = f'=IF({sel}="","",${HC}${39 + i})'
        style(ws[f"W{r}"], size=11, bold=(code == "当日"), align="right", fmt='0"日"', color=(C_RED if code == "当日" else C_INK))
        for c in "RSTUVWX":
            ws[f"{c}{r + 1}"].border = Border(bottom=hair)
    ws.merge_cells("R45:X47")
    ws["R45"] = "確定していたシフトを「休み」に変えた日を、変更した時期ごとに数えたものです。欠勤として数えるのは「当日」のみです。"
    style(ws["R45"], size=8, color=C_MUTED, align="left", valign="top", wrap=True)

    # ---- 2ページ目: 月別サマリー
    ws.row_dimensions[48].height = 14
    ws.row_breaks.append(Break(id=48))
    ws.row_dimensions[49].height = 30
    section(49, "月別の実績")
    ws["B49"].alignment = Alignment(vertical="bottom")
    ws.row_dimensions[50].height = 8
    ws.row_dimensions[51].height = 18
    tbl_cols = [("月", 2, 5), ("出勤日数", 6, 7), ("欠勤日数", 8, 9), ("確定シフト日数", 10, 12), ("出勤率", 13, 15),
                ("勤務時間", 16, 18), ("1日あたり", 19, 20), ("深夜勤務", 21, 22), ("前日休み変更", 23, 24)]
    for title, a, b in tbl_cols:
        ws.merge_cells(f"{L(a)}51:{L(b)}51")
        head_cell(f"{L(a)}51", title, align=("left" if a == 2 else "center"))
    for k in range(1, NSHEETS + 1):
        r = 51 + k
        ws.row_dimensions[r].height = 18
        mw, ma, mh, has = MONTH_WORK[k - 1], MONTH_ABS[k - 1], MONTH_HRS[k - 1], MONTH_HAS[k - 1]
        H = f"N({V(has)})=1"
        vals = [
            (f"={ros}!$AY${21 + k}", None, "left"),
            (f'=IF(${HC}$4="","",IF({H},{V(mw)},"—"))', '0"日"', "right"),
            (f'=IF(${HC}$4="","",IF({H},{V(ma)},"—"))', '0"日"', "right"),
            (f'=IF(${HC}$4="","",IF({H},N({V(mw)})+N({V(ma)}),"—"))', '0"日"', "right"),
            (f'=IF(${HC}$4="","",IF({H},IF(N(J{r})=0,"－",N(F{r})/N(J{r})),"—"))', "0.0%", "right"),
            (f'=IF(${HC}$4="","",IF({H},{V(mh)},"—"))', '0.0"h"', "right"),
            (f'=IF(${HC}$4="","",IF({H},IF(N(F{r})=0,"－",N(P{r})/N(F{r})),"—"))', '0.0"h"', "right"),
            (f'=IF(${HC}$4="","",IF({H},{sum6sel_k("N", k, maxrows, sel)},"—"))', '0.0"h"', "right"),
            (f'=IF(${HC}$4="","",IF({H},SUMIFS({calc_rng(k, "U", maxrows)},{calc_rng(k, "A", maxrows)},{sel},{calc_rng(k, "P", maxrows)},"前日"),"—"))', '0"日"', "right"),
        ]
        for (title, a, b), (formula, fmt, align) in zip(tbl_cols, vals):
            ws.merge_cells(f"{L(a)}{r}:{L(b)}{r}")
            ws[f"{L(a)}{r}"] = formula
            style(ws[f"{L(a)}{r}"], size=10, align=align, fmt=fmt)
            for c in range(a, b + 1):
                ws[f"{L(c)}{r}"].border = Border(bottom=hair)
    r = 58
    ws.row_dimensions[r].height = 18
    totals = [
        ('="合計"', None, "left"),
        (f'=IF(${HC}$4="","",SUM(F52:F57))', '0"日"', "right"),
        (f'=IF(${HC}$4="","",SUM(H52:H57))', '0"日"', "right"),
        (f'=IF(${HC}$4="","",SUM(J52:J57))', '0"日"', "right"),
        (f'=IF(${HC}$4="","",IF(N(J58)=0,"－",N(F58)/N(J58)))', "0.0%", "right"),
        (f'=IF(${HC}$4="","",SUM(P52:P57))', '0.0"h"', "right"),
        (f'=IF(${HC}$4="","",IF(N(F58)=0,"－",N(P58)/N(F58)))', '0.0"h"', "right"),
        (f'=IF(${HC}$4="","",SUM(U52:U57))', '0.0"h"', "right"),
        (f'=IF(${HC}$4="","",SUM(W52:W57))', '0"日"', "right"),
    ]
    for (title, a, b), (formula, fmt, align) in zip(tbl_cols, totals):
        ws.merge_cells(f"{L(a)}{r}:{L(b)}{r}")
        ws[f"{L(a)}{r}"] = formula
        style(ws[f"{L(a)}{r}"], size=10, bold=True, align=align, fmt=fmt)
        for c in range(a, b + 1):
            ws[f"{L(c)}{r}"].border = Border(top=thin, bottom=thin)
    for formula, color in [(f'AND(ISNUMBER(M52),M52>=${HC}$7)', C_GOOD_TXT), (f'AND(ISNUMBER(M52),M52>=${HC}$8)', C_WARN_TXT), ('ISNUMBER(M52)', C_CRIT)]:
        ws.conditional_formatting.add("M52:M58", FormulaRule(formula=[formula], font=Font(name=FONT, bold=True, size=10, color=color), stopIfTrue=True))

    # ---- 欠勤（当日休み変更）の日付一覧
    ws.row_dimensions[59].height = 16
    ws.row_dimensions[60].height = 18
    section(60, "欠勤（当日に休みへ変更）の一覧")
    ws.row_dimensions[61].height = 8
    ws.row_dimensions[62].height = 18
    list_cols = [("日付", 2, 5), ("曜", 6, 6), ("休みへ変更した日時", 7, 11), ("元のシフト時間", 12, 16), ("募集の職種", 17, 24)]
    for title, a, b in list_cols:
        ws.merge_cells(f"{L(a)}62:{L(b)}62")
        head_cell(f"{L(a)}62", title, align=("left" if a in (2, 17) else "center"))

    def lk(colref, i):
        expr = '""'
        for k in range(NSHEETS, 0, -1):
            c = q(S_CALC.format(k))
            expr = f'IFERROR(INDEX({c}!${colref}$2:${colref}${last},MATCH({i},{c}!$W$2:$W${last},0)),{expr})'
        return expr

    for i in range(1, ABS_ROWS + 1):
        r = 62 + i
        ws.row_dimensions[r].height = 17
        date_f = f"={lk('B', i)}" if i > 1 else f'=IF(AND({sel}<>"",N(${HC}$54)=0),"該当なし",{lk("B", i)})'
        cells = [
            (2, 5, date_f, "m/d", "left"),
            (6, 6, f'=IF(ISNUMBER(B{r}),CHOOSE(WEEKDAY(B{r}),"日","月","火","水","木","金","土"),"")', None, "center"),
            (7, 11, f"={lk('V', i)}", "m/d hh:mm", "left"),
            (12, 13, f"={lk('X', i)}", "[h]:mm", "right"),
            (14, 14, f'=IF(ISNUMBER(L{r}),"〜","")', None, "center"),
            (15, 16, f"={lk('Y', i)}", "[h]:mm", "left"),
            (17, 24, f"={lk('R', i)}", None, "left"),
        ]
        for a, b, formula, fmt, align in cells:
            if a != b:
                ws.merge_cells(f"{L(a)}{r}:{L(b)}{r}")
            ws[f"{L(a)}{r}"] = formula
            style(ws[f"{L(a)}{r}"], size=10, align=align, fmt=fmt)
            for c in range(a, b + 1):
                ws[f"{L(c)}{r}"].border = Border(bottom=hair)
    r = 63 + ABS_ROWS
    ws.row_dimensions[r].height = 16
    ws.merge_cells(f"B{r}:X{r}")
    ws[f"B{r}"] = (f'=IF({sel}="","",IF(N(${HC}$54)>{ABS_ROWS},"ほか "&(${HC}$54-{ABS_ROWS})&" 件（表示は最初の{ABS_ROWS}件）。",""))'
                   f'&"日時はシフト管理アプリ上で休みに変更された時刻です。"')
    style(ws[f"B{r}"], size=8, color=C_MUTED, align="left")

    # ---- 注記
    r0 = r + 2
    ws.row_dimensions[r + 1].height = 14
    ws.row_dimensions[r0].height = 18
    section(r0, "集計のきまり")
    ws.row_dimensions[r0 + 1].height = 6
    notes = [
        "出勤日数は確定シフトのあった日数（同じ日の複数区分は1日）。欠勤日数は、確定していたシフトを当日に「休み」へ変更した日数で、前日までの変更は含みません。",
        "出勤率 ＝ 出勤日数 ÷ 確定シフト日数（出勤日数＋欠勤日数）。勤務時間はシフト上の時間（変更後の時間から休憩を除く）。",
        f'="総合判定は出勤率のみによる目安です。確定シフト日数が "&${HC}$9&"日未満の場合は参考値とし、判定を行いません。"',
        f'="欠勤には店側の都合による当日変更が含まれることがあります。面談では本人に事情を確認のうえご利用ください。"&IF({sel}="","","　応募の却下（店側の判断）："&${HC}$55&"件")',
    ]
    for i, t in enumerate(notes):
        rr = r0 + 2 + i
        ws.row_dimensions[rr].height = 26
        ws.merge_cells(f"B{rr}:X{rr}")
        ws[f"B{rr}"] = t
        style(ws[f"B{rr}"], size=8.5, color=C_INK2, align="left", valign="top", wrap=True)

    # ---- 印刷設定（A4縦・幅1ページ・2ページ目に月別以降）
    ws.print_area = f"A1:Y{LAST_ROW}"
    ws.print_title_rows = "1:8"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddFooter.right.text = "&P"
    ws.sheet_view.zoomScale = 100
    return ws


def sum6sel_k(colref, k, maxrows, sel):
    return f"SUMIFS({calc_rng(k, colref, maxrows)},{calc_rng(k, 'A', maxrows)},{sel})"


# ---------------------------------------------------------------- シート: スタッフ一覧
def build_staff_list(wb):
    ws = wb.create_sheet(S_LIST)
    ws.sheet_properties.tabColor = C_BLUE
    ws.sheet_view.showGridLines = False
    ros, sets = q(S_ROSTER), q(S_SET)
    ML = MAXSTAFF + 1

    ws["A1"] = "スタッフ一覧（半年集計）"
    style(ws["A1"], size=16, bold=True)
    ws["A2"] = f'="対象期間　"&{ros}!$AY$29&"　　名前順（先頭の部門記号でまとまります）。フィルターで所属や判定を絞り込めます。"'
    style(ws["A2"], size=9, color=C_INK2)

    # 全体サマリー
    tiles = [("スタッフ数", f"={ros}!$AY$2", '0"名"'), ("評価対象（確定シフトあり）", f"={ros}!$AY$8", '0"名"'),
             ("全体出勤率", f"={ros}!$AY$6", "0.0%"), ("全体欠勤率", f"={ros}!$AY$7", "0.0%"),
             ("総出勤日数", f"={ros}!$AY$3", '#,##0"日"'), ("総欠勤日数", f"={ros}!$AY$4", '#,##0"日"')]
    for i, (label, formula, fmt) in enumerate(tiles):
        c = 1 + i * 2
        a, b = L(c), L(c + 1)
        ws[f"{a}4"] = label
        style(ws[f"{a}4"], size=8, color=C_INK2, bg="FFFFFF")
        ws.merge_cells(f"{a}5:{b}5")
        ws[f"{a}5"] = formula
        style(ws[f"{a}5"], size=16, bold=True, bg="FFFFFF", fmt=fmt, align="left")
        style(ws[f"{b}4"], bg="FFFFFF")
        box_range(ws, f"{a}4:{b}5")
    ws.row_dimensions[5].height = 24

    # 所属別
    ws["A7"] = "所属別"
    style(ws["A7"], size=10, bold=True)
    for j, h in enumerate(["所属", "人数", "出勤日数", "欠勤日数", "確定シフト日数", "出勤率", "欠勤率"]):
        c = ws.cell(row=8, column=1 + j, value=h)
        style(c, size=9, bold=True, color=C_INK, bg="E7E6E1", align="center", border=Border(top=thin, bottom=thin))
    for i in range(len(DEPTS)):
        r = 9 + i
        src = 14 + i
        for j, (colref, fmt) in enumerate([("AY", None), ("AZ", '0"名"'), ("BA", '#,##0"日"'), ("BB", '#,##0"日"'), ("BC", '#,##0"日"'), ("BD", "0.0%"), ("BE", "0.0%")]):
            c = ws.cell(row=r, column=1 + j, value=f"={ros}!${colref}${src}")
            style(c, size=10, fmt=fmt, align=("left" if j == 0 else "right"), border=Border(bottom=hair))

    # ワースト
    ws["J7"] = f'="出勤率 ワースト10（確定シフト日数 "&{sets}!$B$5&"日以上の人が対象）"'
    style(ws["J7"], size=10, bold=True)
    for j, h in enumerate(["名前", "所属", "確定シフト日数", "欠勤日数", "出勤率"]):
        c = ws.cell(row=8, column=10 + j, value=h)
        style(c, size=9, bold=True, color=C_INK, bg="E7E6E1", align="center", border=Border(top=thin, bottom=thin))
    for n in range(1, 11):
        r = 8 + n
        m = f"IFERROR(MATCH({n},{ros}!$Y$2:$Y${ML},0),\"\")"
        ws[f"J{r}"] = f'=IF({m}="","",INDEX({ros}!$J$2:$J${ML},{m}))'
        ws[f"K{r}"] = f'=IF({m}="","",IFERROR(INDEX({ros}!$AY$14:$AY$17,MATCH(INDEX({ros}!$K$2:$K${ML},{m}),{ros}!$AX$14:$AX$17,0)),INDEX({ros}!$K$2:$K${ML},{m})))'
        ws[f"L{r}"] = f'=IF({m}="","",INDEX({ros}!$O$2:$O${ML},{m}))'
        ws[f"M{r}"] = f'=IF({m}="","",INDEX({ros}!$N$2:$N${ML},{m}))'
        ws[f"N{r}"] = f'=IF({m}="","",INDEX({ros}!$P$2:$P${ML},{m}))'
        for c, fmt, al in (("J", None, "left"), ("K", None, "left"), ("L", '0"日"', "right"), ("M", '0"日"', "right"), ("N", "0.0%", "right")):
            style(ws[f"{c}{r}"], size=10, fmt=fmt, align=al, border=Border(bottom=hair))
    ws.conditional_formatting.add("N9:N18", FormulaRule(formula=[f'AND(ISNUMBER(N9),N9<{sets}!$B$4)'], font=Font(name=FONT, bold=True, color=C_CRIT, size=10)))

    # スタッフ別テーブル
    HR = 21
    ws[f"A{HR - 1}"] = "スタッフ別 集計（半年）"
    style(ws[f"A{HR - 1}"], size=10, bold=True)
    headers = ["順位", "名前", "従業員番号", "所属", "資格", "出勤日数", "欠勤日数", "確定シフト日数", "出勤率", "欠勤率", "判定",
               "シフト勤務時間", "平均 h／日", "深夜時間", "却下回数（参考）", "在籍月数"]
    month_hdr_w = [f'="出勤 "&{ros}!$AY${21 + k}' for k in range(1, NSHEETS + 1)]
    month_hdr_a = [f'="欠勤 "&{ros}!$AY${21 + k}' for k in range(1, NSHEETS + 1)]
    for j, h in enumerate(headers + month_hdr_w + month_hdr_a):
        c = ws.cell(row=HR, column=1 + j, value=h)
        style(c, size=9, bold=True, color=C_INK, bg="E7E6E1", align="center", wrap=True, border=Border(top=thin, bottom=thin))
    ws.row_dimensions[HR].height = 32
    widths = [6, 24, 11, 9, 11, 9, 9, 11, 9, 9, 13, 9, 10, 9, 9, 8] + [9] * 12
    for j, w in enumerate(widths):
        ws.column_dimensions[L(1 + j)].width = w

    grade = lambda x: f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"01アルバイト","アルバイト"),"02サブリーダー","サブリーダー"),"03リーダー","リーダー")'
    for n in range(1, MAXSTAFF + 1):
        r = HR + n
        m = f"{ros}!$AV${n + 1}"
        def rv(colref):
            return f"INDEX({ros}!${colref}$2:${colref}${ML},{m})"
        cells = [
            (f'=IF({m}="","",IF({rv("P")}="","－",IF({rv("O")}<{sets}!$B$5,"参考",{rv("W")})))', "0", "center"),
            (f'=IF({m}="","",{rv("J")})', None, "left"),
            (f'=IF({m}="","",{rv("I")})', "0", "right"),
            (f'=IF({m}="","",IFERROR(INDEX({ros}!$AY$14:$AY$17,MATCH({rv("K")},{ros}!$AX$14:$AX$17,0)),{rv("K")}))', None, "left"),
            (f'=IF({m}="","",{grade(rv("L"))})', None, "left"),
            (f'=IF({m}="","",{rv("M")})', '0"日"', "right"),
            (f'=IF({m}="","",{rv("N")})', '0"日"', "right"),
            (f'=IF({m}="","",{rv("O")})', '0"日"', "right"),
            (f'=IF({m}="","",{rv("P")})', "0.0%", "right"),
            (f'=IF({m}="","",{rv("Q")})', "0.0%", "right"),
            (f'=IF({m}="","",IF({rv("P")}="","－",IF({rv("O")}<{sets}!$B$5,"参考値",IF({rv("P")}>={sets}!$B$3,"◎ 良好",IF({rv("P")}>={sets}!$B$4,"△ 注意","✕ 要改善")))))', None, "left"),
            (f'=IF({m}="","",{rv("R")})', '0.0"h"', "right"),
            (f'=IF({m}="","",{rv("T")})', "0.0", "right"),
            (f'=IF({m}="","",{rv("S")})', '0.0"h"', "right"),
            (f'=IF({m}="","",{rv("U")})', '0"回"', "right"),
            (f'=IF({m}="","",{rv("AR")})', '0"ヶ月"', "right"),
        ]
        cells += [(f'=IF({m}="","",IF({ros}!$AZ${22 + k}>0,{rv(MONTH_WORK[k])},""))', "0", "right") for k in range(NSHEETS)]
        cells += [(f'=IF({m}="","",IF({ros}!$AZ${22 + k}>0,{rv(MONTH_ABS[k])},""))', "0", "right") for k in range(NSHEETS)]
        ws.row_dimensions[r].height = 17
        for j, (formula, fmt, al) in enumerate(cells):
            c = ws.cell(row=r, column=1 + j, value=formula)
            style(c, size=10, fmt=fmt, align=al, border=Border(bottom=hair), color=(C_MUTED if j == 14 else C_INK))
    last = HR + MAXSTAFF
    ws.auto_filter.ref = f"A{HR}:{L(16 + 2 * NSHEETS)}{last}"
    ws.freeze_panes = f"C{HR + 1}"
    ws.conditional_formatting.add(f"F{HR + 1}:F{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color=C_BLUE, showValue=True))
    ws.conditional_formatting.add(f"G{HR + 1}:G{last}", FormulaRule(formula=[f"AND(ISNUMBER(G{HR + 1}),G{HR + 1}>0)"], font=Font(name=FONT, bold=True, color=C_CRIT, size=10)))
    for formula, color in [(f'AND(ISNUMBER(I{HR + 1}),I{HR + 1}>={sets}!$B$3)', C_GOOD_TXT), (f'AND(ISNUMBER(I{HR + 1}),I{HR + 1}>={sets}!$B$4)', C_WARN_TXT), (f'ISNUMBER(I{HR + 1})', C_CRIT)]:
        ws.conditional_formatting.add(f"I{HR + 1}:I{last}", FormulaRule(formula=[formula], font=Font(name=FONT, bold=True, size=10, color=color), stopIfTrue=True))
    R1 = HR + 1
    for formula, bg, fg in [(f'AND(ISNUMBER(I{R1}),H{R1}<{sets}!$B$5)', "EEEEEA", C_INK2),
                            (f'AND(ISNUMBER(I{R1}),I{R1}>={sets}!$B$3)', "E3F4E3", C_GOOD_TXT),
                            (f'AND(ISNUMBER(I{R1}),I{R1}>={sets}!$B$4)', "FFF1CC", C_WARN_TXT),
                            (f'ISNUMBER(I{R1})', "F9DEDE", C_CRIT)]:
        ws.conditional_formatting.add(f"K{R1}:K{last}", FormulaRule(formula=[formula], fill=fill(bg), font=Font(name=FONT, bold=True, size=10, color=fg), stopIfTrue=True))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HR}:{HR}"
    ws.print_area = f"A1:P{last}"
    return ws


# ---------------------------------------------------------------- シート: 使い方
def build_howto(wb, maxrows):
    ws = wb.create_sheet(S_HOWTO)
    ws.sheet_properties.tabColor = "EDA100"
    ws.sheet_view.showGridLines = False
    ros = q(S_ROSTER)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 70
    ws["B2"] = "勤務実績ダッシュボード　使い方"
    style(ws["B2"], size=18, bold=True)
    ws["B3"] = "シェアフルシフトの月次シフトCSVを貼り付けるだけで、スタッフごとの出勤日数・出勤率・欠勤率などを集計します。"
    style(ws["B3"], size=10, color=C_INK2)

    r = 5
    ws[f"B{r}"] = "手順"
    style(ws[f"B{r}"], size=12, bold=True)
    steps = [
        ("1", "シェアフルシフトから1ヶ月分のシフトCSVを出力し、Excelで開きます（ダブルクリックで開けます）。"),
        ("2", "CSVの全体（1行目のヘッダーを含む・列はそのまま）をコピーします。Ctrl+A → Ctrl+C。"),
        ("3", "このブックの「CSV_1」シートのA1セルを選択して貼り付けます（Ctrl+V）。2ヶ月目は「CSV_2」、以降「CSV_3」…「CSV_6」へ。順番は古い月から新しい月の順が見やすいです。"),
        ("4", "「ダッシュボード」シートで、スタッフ名をドロップダウンから選ぶ（または氏名を入力する）と、その人の実績が表示されます。印刷はA4縦で、1ページ目に概要とグラフ、2ページ目に月別の実績と欠勤の一覧が入ります。"),
        ("5", "全員の一覧・順位・所属別の集計は「スタッフ一覧」シートで確認できます。判定の基準値や深夜時間帯は「設定」シートで変更できます。"),
        ("★", "貼り直すときは、貼付シートの古いデータをすべて削除（Ctrl+A → Delete）してから貼り付けてください。行数が前より少ない月を上書きすると、古い行が残ってしまいます。"),
    ]
    for i, (n, t) in enumerate(steps):
        rr = r + 1 + i
        ws[f"B{rr}"] = n
        style(ws[f"B{rr}"], size=11, bold=True, color=C_BLUE, align="center")
        ws.merge_cells(f"C{rr}:E{rr}")
        ws[f"C{rr}"] = t
        style(ws[f"C{rr}"], size=10, align="left", wrap=True)
        ws.row_dimensions[rr].height = 30

    r = 13
    ws[f"B{r}"] = "貼付状況"
    style(ws[f"B{r}"], size=12, bold=True)
    for j, h in enumerate(["シート", "月", "貼付行数", "状態"]):
        c = ws.cell(row=r + 1, column=2 + j, value=h)
        style(c, size=9, bold=True, color=C_INK, bg="E7E6E1", align="center", border=Border(top=thin, bottom=thin))
    for k in range(1, NSHEETS + 1):
        rr = r + 1 + k
        ws[f"B{rr}"] = S_CSV.format(k)
        ws[f"C{rr}"] = f"={ros}!$AY${21 + k}"
        ws[f"D{rr}"] = f"={ros}!$AZ${21 + k}"
        ws[f"E{rr}"] = f"={ros}!$BA${21 + k}"
        for c in "BCDE":
            style(ws[f"{c}{rr}"], size=10, border=Border(bottom=hair), align=("right" if c == "D" else "left"))
        ws[f"D{rr}"].number_format = "#,##0"
    ws.conditional_formatting.add(f"E{r + 2}:E{r + 1 + NSHEETS}", FormulaRule(formula=[f'LEFT(E{r + 2},1)="※"'], font=Font(name=FONT, bold=True, color=C_CRIT, size=10)))
    ws.conditional_formatting.add(f"E{r + 2}:E{r + 1 + NSHEETS}", FormulaRule(formula=[f'E{r + 2}="OK"'], font=Font(name=FONT, bold=True, color=C_GOOD_TXT, size=10)))

    r = 22
    ws[f"B{r}"] = "指標の定義"
    style(ws[f"B{r}"], size=12, bold=True)
    defs = [
        ("出勤日数", "ステータスが「確定」系（確定／確定（シフト作成）／確定（アサイン））で、勤務種別が勤務（1）のシフトがある日数。同じ日の複数区分は1日と数えます。"),
        ("欠勤日数", "確定していたシフトが、シフト当日に「休み」（勤務種別4）へ変更された日数。CSVの「更新時間」を日本時間に直し、暦日がシフト日と同じ場合のみ欠勤とします。前日までの変更やシフト日より後の更新は欠勤にも分母にも含めません（ダッシュボードの「休みへの変更タイミング」に表示）。"),
        ("確定シフト日数", "出勤日数 ＋ 欠勤日数。事前に休みへ変更した日は含めません。"),
        ("出勤率／欠勤率", "出勤日数 ÷ 確定シフト日数 ／ 欠勤日数 ÷ 確定シフト日数。"),
        ("勤務時間", "「変更後の開始・終了時間」（無ければ募集時間）から休憩1〜3を引いた、シフト上の時間です。"),
        ("深夜勤務時間", "勤務時間のうち、設定シートの深夜時間帯（既定 22:00〜翌5:00）に重なる時間。"),
        ("却下回数", "応募ステータスが「却下」系の行数。店舗側の採用判断による件数なので、出勤率・判定・順位には用いず参考情報として表示します。"),
        ("総合判定", "出勤率のみを基準にした目安。◎ 良好（設定の基準1以上）／△ 注意（基準2以上）／✕ 要改善（基準2未満）。確定シフト日数が設定の日数未満の人は判定を保留し「参考値」と表示します。"),
        ("順位", "出勤率の高い順。同率は同じ順位（上位の人数＋1）。判定保留の人は順位を出さず「参考」と表示します。"),
    ]
    for i, (k_, t) in enumerate(defs):
        rr = r + 1 + i
        ws[f"B{rr}"] = k_
        style(ws[f"B{rr}"], size=10, bold=True, align="left", valign="top")
        ws.merge_cells(f"C{rr}:E{rr}")
        ws[f"C{rr}"] = t
        style(ws[f"C{rr}"], size=10, align="left", valign="top", wrap=True)
        ws.row_dimensions[rr].height = 30

    r = 34
    ws[f"B{r}"] = "注意事項"
    style(ws[f"B{r}"], size=12, bold=True)
    cautions = [
        f"1ヶ月あたり {maxrows:,} 行、スタッフは半年で {MAXSTAFF} 名まで集計できます。超えた場合は「貼付状況」に警告が出ます。",
        "スタッフは「応募者の従業員番号」で同一人物を判定し、名前は最新月のものを表示します（名前の前後の記号が月によって変わっても同じ人として集計されます）。",
        "CSVの列名（1行目）で列を探すため、列の順番が変わっても動作します。列名が変わった場合は「貼付状況」に警告が出ます。",
        "開いたときに計算に数秒かかることがあります。「ダッシュボード」「スタッフ一覧」以外のシート（設定・名簿・計算1〜6）の数式は編集しないでください。",
        "欠勤（当日休み変更）には、本人都合だけでなく店側都合の当日カットなどが含まれる可能性があります。理由はCSVにないため、必ず本人に事情を確認したうえで評価してください。",
        "「更新時間」は行の最終更新時刻です。シフト日より後に休み行を編集すると「シフト日より後」に分類され欠勤から外れます。その区分が0日でない人は当日の実態を確認してください。",
        "このファイルには個人の勤務情報が含まれます。取り扱いにご注意ください。",
    ]
    for i, t in enumerate(cautions):
        rr = r + 1 + i
        ws.merge_cells(f"B{rr}:E{rr}")
        ws[f"B{rr}"] = "・" + t
        style(ws[f"B{rr}"], size=10, align="left", valign="top", wrap=True)
        ws.row_dimensions[rr].height = 28
    ws.print_area = "A1:E44"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ---------------------------------------------------------------- main
def build(output, csv_paths=(), maxrows=6000, select=None, paste_mode="excel"):
    wb = Workbook()
    wb.remove(wb.active)
    build_howto(wb, maxrows)
    build_dashboard(wb, maxrows, select=select)
    build_staff_list(wb)
    for k in range(1, NSHEETS + 1):
        path = csv_paths[k - 1] if k - 1 < len(csv_paths) else None
        rows = read_csv(path) if path and path != "-" else None
        build_csv_sheet(wb, k, rows, paste_mode)
    build_settings(wb)
    build_roster(wb, maxrows)
    for k in range(1, NSHEETS + 1):
        build_calc_sheet(wb, k, maxrows)

    ML = MAXSTAFF + 1
    wb.defined_names["StaffNames"] = DefinedName(
        "StaffNames", attr_text=f"{q(S_ROSTER)}!$AT$2:INDEX({q(S_ROSTER)}!$AT$2:$AT${ML},MAX(1,N({q(S_ROSTER)}!$AY$2)))")
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.active = wb.sheetnames.index(S_DASH)
    wb.save(output)
    return output


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("output")
    ap.add_argument("--csv", nargs="*", default=[], help="CSV_1, CSV_2 … に貼り付けた状態で生成するCSVファイル（\"-\" で空の月）")
    ap.add_argument("--maxrows", type=int, default=6000, help="1ヶ月あたりの最大行数")
    ap.add_argument("--select", default=None, help="ダッシュボードで初期選択するスタッフ名（検証用）")
    ap.add_argument("--paste-mode", default="excel", choices=["excel", "text"], help="検証用: text はCSVの値をすべて文字列のまま貼り付けた状態を再現")
    a = ap.parse_args()
    print(build(a.output, a.csv, a.maxrows, a.select, a.paste_mode))
