import openpyxl, sys, datetime as dt
wb = openpyxl.load_workbook(sys.argv[1])
ws = wb["祝日・繁忙日"]
r = 17
while ws[f"A{r}"].value is not None:
    r += 1
add = []
D = dt.date
for m in (6, 7):
    for d in range(1, 32):
        try: x = D(2026, m, d)
        except ValueError: continue
        if x.weekday() == 4: add.append((x, "金曜テスト"))
        elif x.weekday() in (1, 2): add.append((x, None))
for x, n in add:
    ws[f"A{r}"] = x; ws[f"A{r}"].number_format = "yyyy/mm/dd"
    ws[f"B{r}"] = n
    r += 1
wb.save(sys.argv[2]); print(len(add), "added up to row", r - 1)
